# -*- coding: utf-8 -*-
"""Pagina unica 'F&I Multibanco | Visao Geral' a partir do HTML de origem.

Replica `source/fmi_multibanco.html` em um slide 16:9 editavel: faixa da
ambicao, tres cartoes de valor e as cinco frentes, todos como shapes nativos
do PowerPoint (nada de imagem colada). Os icones sao desenhados a partir dos
mesmos paths SVG do HTML, entao continuam vetoriais e recoloriveis.

Sistema de coordenadas: o layout e escrito nos pixels do HTML (canvas de
1008x660) e convertido para polegadas na hora de desenhar --- X() usa a
escala horizontal (a pagina abre em largura total, como o `width:100vw` do
HTML) e Y()/F() usam a escala vertical, que define tambem o corpo das
fontes. Icones e circulos usam a escala vertical nos dois eixos, para nao
deformar.

    python3 build_fmi_multibanco.py [saida.pptx]
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from PIL import ImageFont

# ------------------------------------------------------------------ canvas
SW, SH = 13.333, 7.5          # slide 16:9, em polegadas
HW, HH = 1008.0, 660.0        # canvas do HTML, em px
SX = SW / HW                  # px -> polegada (horizontal)
SY = SH / HH                  # px -> polegada (vertical)
FPT = SY * 72.0               # px -> pt (corpo das fontes)

FONT = 'Arial'
LANG = 'pt-BR'

# ------------------------------------------------------------------ tokens
RED_TOP  = RGBColor(0xC4, 0x0A, 0x20)   # gradiente da faixa "Ambicao"
RED_BOT  = RGBColor(0xC1, 0x08, 0x1D)
RED_H1   = RGBColor(0xC7, 0x0A, 0x20)   # segunda linha do titulo
RED_ICON = RGBColor(0xBC, 0x12, 0x27)   # icones dos cartoes de valor
RED_RULE = RGBColor(0xBD, 0x17, 0x30)   # filetes do divisor de secao
RED_NUM  = RGBColor(0xC9, 0x09, 0x1F)   # bolinhas numeradas
RED_FI1  = RGBColor(0xB6, 0x16, 0x2D)   # icone da frente 1
INK      = RGBColor(0x11, 0x11, 0x11)
INK_SOFT = RGBColor(0x22, 0x22, 0x22)
INK_CARD = RGBColor(0x2F, 0x2F, 0x2F)
GREY_TXT = RGBColor(0x44, 0x44, 0x44)
GREY_ICO = RGBColor(0xBC, 0xBC, 0xBC)
DIVIDER  = RGBColor(0xE4, 0xE4, 0xE4)
BG_AMB   = RGBColor(0xF3, 0xF3, 0xF3)
BG_LABEL = RGBColor(0xE8, 0xE8, 0xE8)
BG_CARD  = RGBColor(0xF5, 0xF5, 0xF5)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

# ------------------------------------------------------------------ grade
PAD_L, PAD_R = 34.0, 34.0
CONT_X = PAD_L
CONT_W = HW - PAD_L - PAD_R              # 940
LABEL_W, GAP = 154.0, 10.0

EYEBROW_Y = 14.0
H1_Y      = 45.0
TOP_Y, TOP_H = 138.0, 71.0               # faixa da ambicao
VAL_Y, VAL_H = 221.0, 107.0              # cartoes de valor
SEC_Y     = 354.0                        # divisor "cinco frentes"
FRONTS_Y, FRONTS_H = 392.0, 190.0
FOOTER_Y  = HH - 18.0 - 14.0

CARD_W = (CONT_W - LABEL_W - 3 * GAP) / 3.0          # 252
CARD_X = [CONT_X + LABEL_W + GAP + i * (CARD_W + GAP) for i in range(3)]
FRONT_W = CONT_W / 5.0                               # 188
FRONT_X = [CONT_X + i * FRONT_W for i in range(5)]

_TTF = {False: '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
        True:  '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf'}
_fc = {}


def X(px):
    return px * SX


def Y(px):
    return px * SY


def F(px):
    """corpo da fonte em pt, a partir do tamanho em px do HTML"""
    return round(px * FPT, 1)


def text_px(s, size_px, bold=False):
    """largura do texto em px do HTML (metricas Arial/Liberation Sans)"""
    if bold not in _fc:
        _fc[bold] = ImageFont.truetype(_TTF[bold], 200)
    return _fc[bold].getlength(s) / 200.0 * size_px


# ------------------------------------------------------------------ helpers
def _lang(font):
    rPr = font._rPr
    rPr.set('lang', LANG)
    rPr.set('dirty', '0')


def txt(sl, x, y, w, h, lines, size=14, color=INK, bold=False, lh=1.18,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spc=None, space_after=0):
    """Caixa de texto posicionada em px do HTML.

    lines: str | [(texto, opts)] | [[(texto, opts)], ...] (um item = um paragrafo)
    """
    if isinstance(lines, str):
        lines = [[(lines, {})]]
    elif lines and isinstance(lines[0], tuple):
        lines = [lines]
    box = sl.shapes.add_textbox(Inches(X(x)), Inches(Y(y)), Inches(X(w)), Inches(Y(h)))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, runs in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        psize = runs[0][1].get('size', size) if runs else size
        p.line_spacing = Pt(F(psize) * runs[0][1].get('lh', lh))
        if space_after and i < len(lines) - 1:
            p.space_after = Pt(F(space_after))
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(F(o.get('size', size)))
            f.bold = o.get('bold', bold)
            f.color.rgb = o.get('color', color)
            s = o.get('spc', spc)
            if s:
                f._rPr.set('spc', str(int(round(F(s) * 100))))
            _lang(f)
    return box


def rect(sl, x, y, w, h, fill=None, shape=MSO_SHAPE.RECTANGLE, line=None, lw=0.75):
    sh = sl.shapes.add_shape(shape, Inches(X(x)), Inches(Y(y)),
                             Inches(X(w)), Inches(Y(h)))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return sh


def label(sh, lines, size=14, color=WHITE, bold=False, lh=1.15,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    if isinstance(lines, str):
        lines = [[(lines, {})]]
    elif lines and isinstance(lines[0], tuple):
        lines = [lines]
    tf = sh.text_frame
    tf.vertical_anchor = anchor
    for i, runs in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = Pt(F(size) * lh)
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(F(o.get('size', size)))
            f.bold = o.get('bold', bold)
            f.color.rgb = o.get('color', color)
            _lang(f)
    return sh


def gradient_box(sl, x, y, w, h, top, bottom):
    """faixa com o gradiente vertical do HTML (linear-gradient 180deg)"""
    sh = rect(sl, x, y, w, h, top)
    sh.fill.gradient()
    stops = sh.fill.gradient_stops
    stops[0].color.rgb = top
    stops[0].position = 0.0
    stops[1].color.rgb = bottom
    stops[1].position = 1.0
    sh.fill.gradient_angle = 270.0   # OOXML: 90 graus no sentido horario (cima -> baixo)
    return sh


def vline(sl, x, y1, y2, color, lw=0.75):
    cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(X(x)), Inches(Y(y1)),
                                 Inches(X(x)), Inches(Y(y2)))
    cn.line.color.rgb = color
    cn.line.width = Pt(lw)
    return cn


def hrule(sl, x1, x2, y, color, thick_px=2.0):
    """filete horizontal desenhado como retangulo (espessura em px do HTML)"""
    return rect(sl, x1, y - thick_px / 2.0, x2 - x1, thick_px, color)


# ------------------------------------------------------- icones (viewBox 32)
def _pt(ox, oy, size, a, b):
    """ponto do viewBox 32x32 -> polegadas (escala uniforme)"""
    return Inches(X(ox) + a / 32.0 * size), Inches(Y(oy) + b / 32.0 * size)


def _stroke(sh, color, lw):
    sh.shadow.inherit = False
    sh.fill.background()
    sh.line.color.rgb = color
    sh.line.width = Pt(lw)
    return sh


def _seg(sl, ox, oy, size, p1, p2, color, lw):
    """segmento reto: conector (freeform de bbox nula nao renderiza)"""
    x1, y1 = _pt(ox, oy, size, p1[0], p1[1])
    x2, y2 = _pt(ox, oy, size, p2[0], p2[1])
    cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cn.line.color.rgb = color
    cn.line.width = Pt(lw)
    return cn


def _poly(sl, ox, oy, size, pts, color, lw, close=False):
    if len(pts) == 2 and not close:
        return _seg(sl, ox, oy, size, pts[0], pts[1], color, lw)
    x0, y0 = _pt(ox, oy, size, pts[0][0], pts[0][1])
    b = sl.shapes.build_freeform(x0, y0)
    b.add_line_segments([_pt(ox, oy, size, a, c) for a, c in pts[1:]], close=close)
    return _stroke(b.convert_to_shape(), color, lw)


def _circle(sl, ox, oy, size, cx, cy, r, color, lw):
    x, y = _pt(ox, oy, size, cx - r, cy - r)
    sh = sl.shapes.add_shape(MSO_SHAPE.OVAL, x, y,
                             Inches(2 * r / 32.0 * size), Inches(2 * r / 32.0 * size))
    return _stroke(sh, color, lw)


def _bezier(p0, p1, p2, p3, n=14):
    out = []
    for i in range(1, n + 1):
        t = i / float(n)
        u = 1 - t
        out.append((u * u * u * p0[0] + 3 * u * u * t * p1[0] + 3 * u * t * t * p2[0] + t * t * t * p3[0],
                    u * u * u * p0[1] + 3 * u * u * t * p1[1] + 3 * u * t * t * p2[1] + t * t * t * p3[1]))
    return out


def _arc_up(cx, cy, r, n=12):
    """semicircunferencia com a barriga para cima, da direita para a esquerda"""
    import math
    return [(cx + r * math.cos(math.pi * i / n), cy - r * math.sin(math.pi * i / n))
            for i in range(n + 1)]


def icon_people(sl, ox, oy, size, color, lw):
    _circle(sl, ox, oy, size, 11, 10, 4, color, lw)
    _circle(sl, ox, oy, size, 21, 10, 4, color, lw)
    a = [(4, 25)] + _bezier((4, 25), (4.5, 19), (8, 16), (12, 16)) \
                  + _bezier((12, 16), (16, 16), (19.5, 19), (20, 25))
    b = [(13, 25)] + _bezier((13, 25), (13.5, 20), (16.5, 17.5), (20, 17.5)) \
                   + _bezier((20, 17.5), (23.5, 17.5), (26.5, 20), (27, 25))
    _poly(sl, ox, oy, size, a, color, lw)
    _poly(sl, ox, oy, size, b, color, lw)


def icon_bank(sl, ox, oy, size, color, lw):
    _poly(sl, ox, oy, size, [(5, 12), (16, 5), (27, 12)], color, lw)
    _poly(sl, ox, oy, size, [(7, 13), (25, 13)], color, lw)
    for cx in (9, 15, 21):
        _poly(sl, ox, oy, size, [(cx, 13), (cx, 25)], color, lw)
    _poly(sl, ox, oy, size, [(5, 26), (27, 26)], color, lw)


def icon_car(sl, ox, oy, size, color, lw):
    body = [(4, 19), (7, 12), (25, 12), (28, 19), (28, 25), (25, 25)] \
        + _arc_up(21, 25, 4)[1:] + [(15, 25)] + _arc_up(11, 25, 4)[1:] + [(4, 25)]
    _poly(sl, ox, oy, size, body, color, lw, close=True)
    _poly(sl, ox, oy, size, [(8, 12), (11, 8), (21, 8), (24, 12)], color, lw)
    _circle(sl, ox, oy, size, 11, 24, 2.5, color, lw)
    _circle(sl, ox, oy, size, 21, 24, 2.5, color, lw)


def icon_tech(sl, ox, oy, size, color, lw):
    _circle(sl, ox, oy, size, 16, 16, 5, color, lw)
    _circle(sl, ox, oy, size, 16, 16, 10, color, lw)
    for seg in ((16, 3, 16, 7), (16, 25, 16, 29), (3, 16, 7, 16), (25, 16, 29, 16),
                (7, 7, 10, 10), (22, 22, 25, 25), (25, 7, 22, 10), (10, 22, 7, 25)):
        _poly(sl, ox, oy, size, [(seg[0], seg[1]), (seg[2], seg[3])], color, lw)


def icon_screen(sl, ox, oy, size, color, lw):
    _poly(sl, ox, oy, size, [(5, 5), (27, 5), (27, 21), (5, 21)], color, lw, close=True)
    _poly(sl, ox, oy, size, [(12, 27), (20, 27)], color, lw)
    _poly(sl, ox, oy, size, [(16, 21), (16, 27)], color, lw)


def icon_network(sl, ox, oy, size, color, lw):
    _circle(sl, ox, oy, size, 8, 16, 3, color, lw)
    _circle(sl, ox, oy, size, 23, 8, 3, color, lw)
    _circle(sl, ox, oy, size, 23, 24, 3, color, lw)
    _poly(sl, ox, oy, size, [(11, 15), (20, 9)], color, lw)
    _poly(sl, ox, oy, size, [(11, 17), (20, 23)], color, lw)


def draw_icon(sl, kind, x, y, size_px, color):
    """size_px = lado do icone em px do HTML; traco = stroke-width 2 do SVG"""
    size = Y(size_px)                      # escala uniforme (nao deforma)
    lw = max(0.75, 2.0 / 32.0 * size_px * FPT)
    {'people': icon_people, 'bank': icon_bank, 'car': icon_car,
     'tech': icon_tech, 'screen': icon_screen, 'network': icon_network}[kind](
        sl, x, y, size, color, lw)


# ------------------------------------------------------------------ conteudo
EYEBROW = 'F&I MULTIBANCO | VISÃO GERAL'
H1 = [('Entregar a plataforma de F&I multibanco depende de cinco frentes,', INK),
      ('iniciadas por um piloto com lojistas', RED_H1)]
AMBICAO = ['Plataforma de F&I multibanco e ponta a ponta que faz o lojista vender mais,',
           'com simulação e envio de ficha em um só fluxo']
BENEFITS = [
    ('people', 'Lojista e seu cliente',
     'Mais aprovação, menos retrabalho e resposta em minutos'),
    ('car', 'Localiza',
     'Mais penetração de F&I e relação mais forte com a rede de lojistas'),
    ('bank', 'Banco',
     'Originação qualificada com custo de aquisição menor'),
]
SECTION = 'cinco frentes, escaladas em ondas'
FRONTS = [
    ('bank', ['Bancos'],
     'Integrar em ondas, dos bancos menores aos maiores, à medida que a '
     'proposta de valor se confirma'),
    ('tech', ['Tecnologia'],
     'APIs para simulação via Corban do lojista e envio da ficha a múltiplos '
     'bancos em um único disparo'),
    ('screen', ['Interface', 'da plataforma'],
     'Melhoria contínua guiada pelo uso do vendedor, com ciclos curtos de release'),
    ('people', ['Lojistas'],
     'Piloto em pequena escala para provar a proposta de valor e, confirmada, '
     'escalar por onda'),
    ('network', ['Ecossistema'],
     'Novos serviços e integração com Valorização depois que o núcleo estiver estável'),
]
FOOTNOTE = 'Fonte: análise Bain. Preliminar, para discussão'
PAGENO = '2'


def page(prs):
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    # ---- eyebrow
    txt(sl, CONT_X, EYEBROW_Y, CONT_W, 16, EYEBROW, size=11, bold=True,
        color=GREY_TXT, spc=4, lh=1.2)

    # ---- titulo (duas linhas, a segunda em vermelho)
    txt(sl, CONT_X, H1_Y, 920, 70,
        [[(H1[0][0], {'color': H1[0][1]})], [(H1[1][0], {'color': H1[1][1]})]],
        size=31, bold=True, lh=1.05)

    # ---- ambicao
    amb = gradient_box(sl, CONT_X, TOP_Y, LABEL_W, TOP_H, RED_TOP, RED_BOT)
    amb.text_frame.margin_left = Inches(X(18))
    label(amb, 'Ambição', size=19, bold=True, color=WHITE)
    rect(sl, CONT_X + LABEL_W + GAP, TOP_Y, CONT_W - LABEL_W - GAP, TOP_H, BG_AMB)
    txt(sl, CONT_X + LABEL_W + GAP + 17, TOP_Y + 12, CONT_W - LABEL_W - GAP - 34,
        TOP_H - 24, [[(AMBICAO[0], {})], [(AMBICAO[1], {})]],
        size=20, bold=True, lh=1.17)

    # ---- valor para...
    lbl = rect(sl, CONT_X, VAL_Y, LABEL_W, VAL_H, BG_LABEL)
    lbl.text_frame.margin_left = Inches(X(18))
    label(lbl, 'Valor para...', size=18, bold=True, color=INK)
    for i, (kind, head, body) in enumerate(BENEFITS):
        cx = CARD_X[i]
        rect(sl, cx, VAL_Y, CARD_W, VAL_H, BG_CARD)
        draw_icon(sl, kind, cx + 16, VAL_Y + 14, 30, RED_ICON)
        txt(sl, cx + 60, VAL_Y + 13, CARD_W - 76, VAL_H - 25,
            [[(head, {'size': 18, 'bold': True, 'color': INK, 'lh': 1.15})],
             [(body, {'size': 16, 'color': INK_CARD, 'lh': 1.2})]],
            space_after=2)

    # ---- divisor de secao
    sec_mid = SEC_Y + 11
    hrule(sl, CONT_X, CONT_X + 104, sec_mid, RED_RULE)
    ico_x = CONT_X + 104 + 8
    rect(sl, ico_x, sec_mid - 9, 27, 18, None, MSO_SHAPE.OVAL, GREY_ICO, 1.5)
    sec_txt_x = ico_x + 27 + 8
    hrule(sl, ico_x + 18, sec_txt_x, sec_mid, RED_RULE)
    sec_w = text_px(SECTION, 19, True)
    txt(sl, sec_txt_x, SEC_Y, sec_w + 12, 24, SECTION, size=19, bold=True,
        color=INK, lh=1.15)
    hrule(sl, sec_txt_x + sec_w + 8, CONT_X + CONT_W, sec_mid, RED_RULE)

    # ---- cinco frentes
    for i, (kind, head, body) in enumerate(FRONTS):
        fx = FRONT_X[i]
        if i < 4:
            vline(sl, fx + FRONT_W, FRONTS_Y, FRONTS_Y + FRONTS_H, DIVIDER)
        draw_icon(sl, kind, fx + 24, FRONTS_Y + 10, 35,
                  RED_FI1 if i == 0 else INK)
        num = rect(sl, fx, FRONTS_Y - 2, 29 * SY / SX, 29, RED_NUM, MSO_SHAPE.OVAL)
        label(num, str(i + 1), size=16, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER)
        h3_h = 17 * 1.05 * len(head)
        txt(sl, fx + 16, FRONTS_Y + 56, FRONT_W - 33, h3_h + 4,
            [[(l, {})] for l in head], size=17, bold=True, color=INK, lh=1.05)
        txt(sl, fx + 16, FRONTS_Y + 56 + h3_h + 11, FRONT_W - 33,
            FRONTS_H - 56 - h3_h - 11, body, size=14, color=INK_SOFT, lh=1.18)

    # ---- rodape
    txt(sl, CONT_X, FOOTER_Y, 600, 16, FOOTNOTE, size=11, color=GREY_TXT, lh=1.2)
    txt(sl, CONT_X + CONT_W - 100, FOOTER_Y - 1, 100, 18, PAGENO, size=13,
        color=GREY_TXT, align=PP_ALIGN.RIGHT, lh=1.2)
    return sl


def main(dst):
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    page(prs)
    cp = prs.core_properties
    cp.title = 'F&I Multibanco | Visão Geral'
    cp.author = 'ConsultingInnovation@bain.com'
    cp.last_modified_by = 'ConsultingInnovation@bain.com'
    prs.save(dst)
    print('ok %s (%d pagina)' % (dst, len(prs.slides._sldIdLst)))


if __name__ == '__main__':
    import sys
    main(sys.argv[1] if len(sys.argv) > 1 else '../output/fmi_multibanco.pptx')
