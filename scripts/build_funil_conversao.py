# -*- coding: utf-8 -*-
"""Pagina 'Funil lead -> venda por canal' (CV, Liza e Central, jan-ago/26).

Reproduz o rascunho de tres paineis (um por canal) com o funil completo do lead
a venda, mes a mes, mais a coluna de meta. As taxas sao recalculadas a partir
dos volumes brutos -- nada e transcrito de valores ja arredondados -- e a linha
de conversao final recebe heatmap contra a meta de 0,94%.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- tokens
RED    = RGBColor(0xCC, 0x00, 0x00)   # Bain red
RUBY   = RGBColor(0x99, 0x00, 0x00)   # negativo
FOREST = RGBColor(0x10, 0x4C, 0x3E)   # positivo
BLACK  = RGBColor(0x00, 0x00, 0x00)
GREY1  = RGBColor(0x33, 0x33, 0x33)
GREY2  = RGBColor(0x5C, 0x5C, 0x5C)
GREY3  = RGBColor(0x85, 0x85, 0x85)
GREY4  = RGBColor(0xB4, 0xB4, 0xB4)
GREY5  = RGBColor(0xDC, 0xDC, 0xDC)
BAND   = RGBColor(0xF4, 0xF4, 0xF4)   # zebra das linhas de taxa
METABG = RGBColor(0xEE, 0xEE, 0xEE)   # coluna de meta
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
FONT   = "Arial"
LANG   = "pt-BR"

# ---------------------------------------------------------------- grid (in)
SW, SH = 13.333, 7.5
MARGIN_L, MARGIN_R = 0.30, 13.03
TITLE_Y, TITLE_H = 0.17, 0.62
RULE_Y = 0.90

HEAD_Y, HEAD_H = 0.99, 0.40          # barra com o nome do canal
READ_Y, READ_H = 1.43, 0.32          # leitura de uma linha por canal
MON_Y,  MON_H  = 1.80, 0.26          # cabecalho de meses
BODY_Y = 2.09
ROW_H  = 0.428
NOTE_Y = 7.02

BAR_X,  BAR_W  = MARGIN_L, 0.20      # bracket vertical da etapa do funil
LBL_X,  LBL_W  = 0.56, 2.02          # rotulo da metrica
PANEL_X, PANEL_GAP = 2.70, 0.14
PANEL_W = (MARGIN_R - PANEL_X - 2 * PANEL_GAP) / 3.0
MON_W   = 0.3525                     # 8 meses
META_GAP, META_W = 0.07, 0.44

MESES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago"]
META_CONV = 0.94                     # meta de conversao lead -> venda (%)

# ---------------------------------------------------------------- dados
# volumes brutos por mes (jan a ago/26); None = canal ainda nao operava
N = None
CANAIS = [
    dict(
        nome="CV (sem Liza e sem Central)",
        meta_leads=400000, meta_leads_lbl="400k",
        leads=[454314, 351045, 387905, 426871, 435764, 336078, 324671, 321406],
        cep=[N, N, N, 0.46, 0.48, 0.47, 0.47, 0.39],
        env=[16048, 13734, 15604, 16497, 17662, 19035, 16326, 15215],
        apr=[5711, 4703, 5161, 5203, 5161, 5256, 4648, 4619],
        fat=[1933, 1553, 1803, 1654, 1575, 1646, 1527, 1377],
        ven=[3416, 2674, 3183, 2816, 2742, 2669, 2480, 2144],
        leitura="Leads −29% (454k→321k) e conversão −11%; eficiência do funil estável",
    ),
    dict(
        nome="Liza",
        meta_leads=50000, meta_leads_lbl="50k",
        leads=[N, N, N, 8117, 13735, 14688, 31018, 58349],
        cep=[N, N, N, 0.54, 0.55, 0.53, 0.49, 0.47],
        env=[N, N, N, 584, 890, 1140, 1557, 2684],
        apr=[N, N, N, 141, 236, 285, 466, 809],
        fat=[N, N, N, 50, 81, 85, 168, 248],
        ven=[N, N, N, 13, 10, 12, 26, 20],
        leitura="Leads 7,2x desde abr/26, mas conversão −79%; vendas < fichas faturadas¹",
    ),
    dict(
        nome="Central",
        meta_leads=20000, meta_leads_lbl="20k",
        leads=[14502, 12085, 16248, 14333, 13311, 10660, 15577, 9060],
        cep=[N, N, N, 0.64, 0.73, 0.68, 0.54, 0.55],
        env=[134, 137, 167, 175, 175, 188, 719, 397],
        apr=[69, 66, 82, 71, 78, 91, 543, 320],
        fat=[29, 35, 43, 28, 29, 32, 88, 68],
        ven=[116, 101, 136, 98, 101, 83, 210, 107],
        leitura="Único acima da meta: conversão +48% e aprovação de 51% para 81%",
    ),
]

# etapas do funil: (nome do bracket, cor, [(rotulo, chave de calculo)])
ETAPAS = [
    ("LEADS", GREY1, [
        ("# de leads",                        "leads"),
        ("% leads sem info de CEP",           "cep"),
    ]),
    ("FICHAS", GREY2, [
        ("# fichas enviadas",                 "env"),
        ("% fichas enviadas / leads",         "env_leads"),
    ]),
    ("APROVAÇÃO", GREY2, [
        ("# fichas aprovadas",                "apr"),
        ("% aprovadas / enviadas",            "apr_env"),
    ]),
    ("FATURAMENTO", GREY2, [
        ("# fichas faturadas",                "fat"),
        ("% faturadas / aprovadas",           "fat_apr"),
        ("% fichas faturadas / leads",        "fat_leads"),
    ]),
    ("VENDA", RED, [
        ("# de vendas",                       "ven"),
        ("% de vendas / leads",               "ven_leads"),
    ]),
]

TITLE = ("Conversão lead→venda caiu ~23% no consolidado (0,75%→0,58%) e segue ~38% abaixo da meta "
         "de 0,94%: só a Central avança (+48%), com Liza a −79% e CV a −11%")

NOTE = ("Nota: CV exclui Liza e Central. Taxas recalculadas sobre os volumes brutos; % fichas enviadas, "
        "% fichas faturadas / leads e % de vendas são sobre o total de leads do mês, % aprovadas sobre "
        "enviadas e % faturadas sobre aprovadas. Meta de conversão de 0,94% aplicada sobre a meta mensal de "
        "leads de cada canal (CV 400k, Liza 50k, Central 20k). (1) Em ago/26 a Liza registra 248 fichas "
        "faturadas e apenas 20 vendas, inversão frente ao CV, o que sugere lacuna de atribuição da venda ao "
        "canal. % leads sem info de CEP não disponível para jan-mar/26; Liza inicia em abr/26. "
        "Fonte: base de leads, base de fichas e base de vendas faturadas")

# ---------------------------------------------------------------- helpers
def _lang(run):
    run.font._rPr.set('lang', LANG)


def txt(slide, x, y, w, h, parts, size=7, color=GREY1, bold=False,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, wrap=True, line=None):
    """parts: str ou lista de (trecho, {bold, color, size})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if line:
        p.line_spacing = line
    if isinstance(parts, str):
        parts = [(parts, {})]
    for t, o in parts:
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = FONT
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = o.get("color", color)
        _lang(r)
    return tb


def rect(slide, x, y, w, h, fill=None, line_color=None, line_w=0.75,
         shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line_color is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_color
        s.line.width = Pt(line_w)
    s.text_frame.word_wrap = False
    return s


def hline(slide, x1, y, x2, color=GREY4, w=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y),
                                   Inches(x2), Inches(y))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    return c


def rot_label(slide, cx, cy, length, thick, text, color, size=6):
    """rotulo girado 270 graus, centrado em (cx, cy) e ocupando thick x length"""
    tb = txt(slide, cx - length / 2, cy - thick / 2, length, thick, text,
             size=size, color=color, bold=True, wrap=False)
    tb.rotation = 270
    return tb


def mix(a, b, t):
    t = max(0.0, min(1.0, t))
    return RGBColor(*[int(round(x + (y - x) * t)) for x, y in
                      ((a[0], b[0]), (a[1], b[1]), (a[2], b[2]))])


# ---------------------------------------------------------------- formatacao
def fmt_milhar(v):
    return "{:,}".format(int(round(v))).replace(",", ".")


def fmt_num_linha(vals):
    """escolhe uma unica unidade para toda a linha, a partir do maior valor"""
    pico = max([v for v in vals if v is not None] or [0])
    if pico >= 100000:
        return lambda v: "%dk" % round(v / 1000.0)
    if pico >= 10000:
        return lambda v: ("%.1fk" % (v / 1000.0)).replace(".", ",")
    return fmt_milhar


def fmt_pct(v, dec):
    if v is None:
        return "–"
    return ("%.*f%%" % (dec, v)).replace(".", ",")


def ratio(num, den):
    if num is None or den in (None, 0):
        return None
    return num / float(den) * 100.0


def serie(canal, chave):
    """retorna (valores, formatador) para uma linha do funil"""
    L, E, A, F, V = (canal[k] for k in ("leads", "env", "apr", "fat", "ven"))
    if chave in ("leads", "env", "apr", "fat", "ven"):
        return canal[chave], fmt_num_linha(canal[chave])
    if chave == "cep":
        return [None if c is None else c * 100 for c in canal["cep"]], lambda v: fmt_pct(v, 0)
    if chave == "env_leads":
        return [ratio(e, l) for e, l in zip(E, L)], lambda v: fmt_pct(v, 1)
    if chave == "apr_env":
        return [ratio(a, e) for a, e in zip(A, E)], lambda v: fmt_pct(v, 0)
    if chave == "fat_apr":
        return [ratio(f, a) for f, a in zip(F, A)], lambda v: fmt_pct(v, 0)
    if chave == "fat_leads":
        return [ratio(f, l) for f, l in zip(F, L)], lambda v: fmt_pct(v, 2)
    if chave == "ven_leads":
        return [ratio(v, l) for v, l in zip(V, L)], lambda v: fmt_pct(v, 2)
    raise KeyError(chave)


def meta_cell(canal, chave):
    """valor da coluna Meta; None = celula vazia"""
    if chave == "leads":
        return canal["meta_leads_lbl"]
    if chave == "ven":
        return fmt_milhar(canal["meta_leads"] * META_CONV / 100.0)
    if chave == "ven_leads":
        return fmt_pct(META_CONV, 2)
    return None


def heat(v):
    """fundo e cor do texto da linha de conversao, contra a meta de 0,94%"""
    if v is None:
        return None, GREY4, False
    r = v / META_CONV
    if r >= 1.0:
        return mix((0xE4, 0xEF, 0xEA), (0xB6, 0xD6, 0xC7), (r - 1.0) / 0.5), FOREST, True
    fill = mix((0xFF, 0xFF, 0xFF), (0xEE, 0xB3, 0xB3), 1.0 - r)
    return fill, (RUBY if r < 0.45 else GREY1), r < 0.45


def delta_conv(canal):
    """variacao da conversao entre o primeiro e o ultimo mes com dados"""
    s, _ = serie(canal, "ven_leads")
    vals = [x for x in s if x is not None]
    return (vals[-1] / vals[0] - 1.0) * 100.0


# ---------------------------------------------------------------- build
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# titulo + regua vermelha
txt(slide, MARGIN_L, TITLE_Y, 10.85, TITLE_H, TITLE, size=18, color=GREY1,
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.06)
txt(slide, 11.25, TITLE_Y + 0.36, 1.78, 0.22, "/ P R E L I M I N A R", size=8.5,
    color=GREY3, align=PP_ALIGN.RIGHT, wrap=False)
hline(slide, MARGIN_L, RULE_Y, MARGIN_R, RED, 2.25)

# cabecalho da coluna de rotulos
txt(slide, BAR_X, HEAD_Y, LBL_X + LBL_W - BAR_X, HEAD_H,
    [("Funil lead → venda", {"bold": True, "size": 10.5, "color": GREY1})],
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
txt(slide, BAR_X, READ_Y, LBL_X + LBL_W - BAR_X, READ_H,
    "Volumes e taxas por mês, jan a ago/26", size=6.5, color=GREY3,
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
txt(slide, BAR_X, MON_Y, LBL_X + LBL_W - BAR_X, MON_H, "2026", size=6.5,
    color=GREY4, bold=True, align=PP_ALIGN.LEFT)

# linhas do funil: brackets de etapa + rotulos
linhas = [(lbl, key, cor) for _, cor, itens in ETAPAS for lbl, key in itens]
y = BODY_Y
for nome_etapa, cor, itens in ETAPAS:
    gh = ROW_H * len(itens)
    rect(slide, BAR_X, y + 0.02, BAR_W, gh - 0.04, fill=cor)
    rot_label(slide, BAR_X + BAR_W / 2, y + gh / 2, gh - 0.06, BAR_W,
              nome_etapa, WHITE, size=5.5 if len(nome_etapa) > 9 else 6)
    for i, (lbl, key) in enumerate(itens):
        ry = y + i * ROW_H
        destaque = key in ("ven_leads", "fat_leads")
        txt(slide, LBL_X, ry, LBL_W, ROW_H, lbl, size=7.5,
            color=GREY1 if lbl.startswith("#") or destaque else GREY2,
            bold=destaque or lbl.startswith("#"),
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    y += gh

BODY_H = ROW_H * len(linhas)

# paineis, um por canal
for ci, canal in enumerate(CANAIS):
    px = PANEL_X + ci * (PANEL_W + PANEL_GAP)
    meta_x = px + PANEL_W - META_W

    # barra do canal + chip de variacao da conversao
    rect(slide, px, HEAD_Y, PANEL_W, HEAD_H, fill=GREY1)
    txt(slide, px + 0.12, HEAD_Y, PANEL_W - 0.95, HEAD_H, canal["nome"], size=10,
        color=WHITE, bold=True, align=PP_ALIGN.LEFT, wrap=False)
    d = delta_conv(canal)
    rect(slide, px + PANEL_W - 0.80, HEAD_Y + 0.07, 0.68, HEAD_H - 0.14, fill=WHITE)
    txt(slide, px + PANEL_W - 0.80, HEAD_Y + 0.07, 0.68, HEAD_H - 0.14,
        "%+d%%" % round(d), size=9, color=FOREST if d > 0 else RUBY, bold=True)

    txt(slide, px, READ_Y, PANEL_W, READ_H, canal["leitura"], size=6.5,
        color=GREY2, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, line=1.05)

    # fundo da coluna de meta + cabecalho de meses
    rect(slide, meta_x, MON_Y, META_W, MON_H + BODY_H, fill=METABG)
    for mi, mes in enumerate(MESES):
        txt(slide, px + mi * MON_W, MON_Y, MON_W, MON_H, mes, size=6.5,
            color=GREY2, bold=True)
    txt(slide, meta_x, MON_Y, META_W, MON_H, "Meta", size=6.5, color=RED, bold=True)
    hline(slide, px, MON_Y + MON_H, px + PANEL_W, GREY1, 1.0)

    # celulas
    for ri, (lbl, key, _) in enumerate(linhas):
        ry = BODY_Y + ri * ROW_H
        vals, fmt = serie(canal, key)
        is_pct = lbl.startswith("%")
        chave_forte = key in ("ven_leads", "fat_leads")

        if key == "ven_leads":
            pass                                  # heatmap por celula, abaixo
        elif chave_forte:
            rect(slide, px, ry, PANEL_W - META_W - META_GAP, ROW_H, fill=BAND)
        elif is_pct:
            rect(slide, px, ry, PANEL_W - META_W - META_GAP, ROW_H,
                 fill=RGBColor(0xFA, 0xFA, 0xFA))

        for mi, v in enumerate(vals):
            cx = px + mi * MON_W
            cor, bold = (GREY1, True) if not is_pct else (GREY2, False)
            if chave_forte:
                cor, bold = GREY1, True
            if key == "ven_leads":
                fill, cor, bold = heat(v)
                if fill is not None:
                    rect(slide, cx + 0.005, ry + 0.03, MON_W - 0.01, ROW_H - 0.06,
                         fill=fill)
            if v is None:
                txt(slide, cx, ry, MON_W, ROW_H, "–", size=7, color=GREY4)
            else:
                txt(slide, cx, ry, MON_W, ROW_H, fmt(v), size=7 if not chave_forte else 7.5,
                    color=cor, bold=bold)

        mv = meta_cell(canal, key)
        if mv is not None:
            txt(slide, meta_x, ry, META_W, ROW_H, mv, size=7.5, color=RED, bold=True)

        if ri < len(linhas) - 1:
            hline(slide, px, ry + ROW_H, px + PANEL_W - META_W - META_GAP, GREY5, 0.5)

    # moldura: conversao final e limites do painel
    y_ven = BODY_Y + (len(linhas) - 1) * ROW_H
    rect(slide, px, y_ven, PANEL_W - META_W - META_GAP, ROW_H, fill=None,
         line_color=RED, line_w=1.0)
    hline(slide, px, BODY_Y + BODY_H, px + PANEL_W, GREY1, 1.0)

# nota / fonte
txt(slide, MARGIN_L, NOTE_Y, MARGIN_R - MARGIN_L, 0.40, NOTE, size=6,
    color=GREY3, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.18)

OUT = "/home/user/Gerais/output/funil_conversao_canais.pptx"
prs.save(OUT)
print("saved", OUT)
