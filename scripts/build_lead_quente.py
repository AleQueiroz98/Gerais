# -*- coding: utf-8 -*-
"""Lead quente: os quatro passos para converter (L2S).

Uma pagina 16:9 que replica, elemento por elemento, o painel HTML
`source/lead_quente.html`:

  - quatro cartoes numerados (comunicar, priorizar, avisar, realocar), cada um
    com icone de linha, titulo, tres bullets e o bloco dono/prazo;
  - chevrons finos marcando a sequencia entre os passos;
  - faixa de merito no rodape, com a medalha e a mensagem-chave da pagina.

Tudo e desenhado do zero com python-pptx e a geometria vem das mesmas
coordenadas do HTML (1280 x 720 CSS px, 96 px por polegada). Os icones vem da
familia `*_linha` de `icons_bain.py`, convertida dos mesmos paths SVG do
painel, entao o deck continua editavel e nao depende de imagem externa.

    python3 scripts/build_lead_quente.py [destino.pptx]
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

from icons_bain import icon

# ------------------------------------------------------------------ tokens
RED    = RGBColor(0xCC, 0x00, 0x00)   # Bain red (o painel usa #c8102e)
ROSE   = RGBColor(0xD9, 0x6C, 0x7E)   # borda da faixa de merito
BLACK  = RGBColor(0x11, 0x11, 0x11)
BODY   = RGBColor(0x1D, 0x1D, 0x1D)
GREY1  = RGBColor(0x33, 0x33, 0x33)
GREY2  = RGBColor(0x47, 0x47, 0x47)
GREY3  = RGBColor(0x55, 0x55, 0x55)
GREY4  = RGBColor(0x6B, 0x6B, 0x6B)
CHEV   = RGBColor(0x22, 0x22, 0x22)
CARD_BG = RGBColor(0xF7, 0xF7, 0xF8)
CARD_LN = RGBColor(0xEC, 0xEC, 0xEC)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
FONT   = 'Arial'
LANG   = 'pt-BR'

SW, SH = 13.333, 7.5          # polegadas
DPI    = 96.0                 # o painel de origem e 1280 x 720 CSS px


def I(px):
    """px do painel HTML -> polegadas do slide"""
    return px / DPI


def pt(px):
    """px do painel HTML -> pontos de fonte"""
    return px * 0.75


# ------------------------------------------------------------------ grade
MX, CW = 18, 1244             # padding lateral do painel

EYE_Y = 10                    # eyebrow "L2S | LEAD QUENTE"
BAR_Y, BAR_W, BAR_H = 34, 35, 4
TITLE_Y, TITLE_H = 46, 60
SUB_Y, SUB_H = 114, 42

GAP = 30
CARD_W = (CW - 3 * GAP) / 4.0
CARD_Y, CARD_H = 172, 368
CARD_X = [MX + i * (CARD_W + GAP) for i in range(4)]

BUBBLE_D = 36
ICON_D = 58
CHEV_CY = CARD_Y + 53

BAN_Y, BAN_H = 552, 45
MEDAL_D = 52
FOOT_Y = 607


# ------------------------------------------------------------------ helpers
def _lang(font):
    rPr = font._rPr
    rPr.set('lang', LANG)
    rPr.set('dirty', '0')


def _spc(font, hundredths):
    font._rPr.set('spc', str(int(hundredths)))


def flat(sh):
    """remove a sombra herdada do tema (p:style -> effectRef)"""
    sp = sh._element
    st = sp.find(qn('p:style'))
    if st is not None:
        sp.remove(st)
    spPr = sp.spPr
    for el in spPr.findall(qn('a:effectLst')):
        spPr.remove(el)
    spPr.append(spPr.makeelement(qn('a:effectLst'), {}))
    sh.shadow.inherit = False
    return sh


def _nobullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum', 'a:buFont'):
        el = pPr.find(qn(tag))
        if el is not None:
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn('a:buNone'), {}))


def _dot(p, mar=152400):
    """bullet redondo com recuo pendurado, como o <ul> do painel"""
    pPr = p._p.get_or_add_pPr()
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum', 'a:buFont'):
        el = pPr.find(qn(tag))
        if el is not None:
            pPr.remove(el)
    pPr.set('marL', str(mar))
    pPr.set('indent', str(-mar))
    pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': FONT}))
    pPr.append(pPr.makeelement(qn('a:buChar'), {'char': u'•'}))


def txt(sl, x, y, w, h, lines, size=11, color=BODY, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spc=None, lsp=1.10,
        space_after=0, bullet=False, name=None):
    """lines: str | [(texto, opts)] | [[(texto, opts)], ...] (paragrafos)"""
    if isinstance(lines, str):
        lines = [[(lines, {})]]
    elif lines and isinstance(lines[0], tuple):
        lines = [lines]
    box = sl.shapes.add_textbox(Inches(I(x)), Inches(I(y)),
                                Inches(I(w)), Inches(I(h)))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, runs in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = lsp
        p.space_before = Pt(0)
        p.space_after = Pt(space_after)
        _dot(p) if bullet else _nobullet(p)
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.color.rgb = o.get('color', color)
            if o.get('spc', spc):
                _spc(f, o.get('spc', spc))
            _lang(f)
    return box


def rect(sl, x, y, w, h, fill=None, line=None, lw=0.75,
         shape=MSO_SHAPE.RECTANGLE, adj=None, name=None):
    sh = sl.shapes.add_shape(shape, Inches(I(x)), Inches(I(y)),
                             Inches(I(w)), Inches(I(h)))
    flat(sh)
    if name:
        sh.name = name
    if adj is not None:
        sh.adjustments[0] = adj
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return sh


def label(sh, text, size=9.5, color=WHITE, bold=True,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf = sh.text_frame
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    _nobullet(p)
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    _lang(r.font)
    return sh


def gradient_card(sl, x, y, w, h, name=None):
    """cartao com o degrade sutil do painel (#f7f7f8 -> branco em 28%)"""
    sh = rect(sl, x, y, w, h, WHITE, CARD_LN, 0.75, name=name)
    sh.fill.gradient()
    sh.fill.gradient_angle = 270.0           # de cima para baixo
    stops = sh.fill.gradient_stops
    stops[0].color.rgb = CARD_BG
    stops[0].position = 0.0
    stops[1].color.rgb = WHITE
    stops[1].position = 0.28
    return sh


def hline(sl, x1, y, x2, color=GREY4, lw=0.75):
    cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(I(x1)),
                                 Inches(I(y)), Inches(I(x2)), Inches(I(y)))
    cn.line.color.rgb = color
    cn.line.width = Pt(lw)
    return cn


# ------------------------------------------------------- formas da pagina
def stroke(sh, w, color):
    flat(sh)
    sh.fill.background()
    sh.line.color.rgb = color
    sh.line.width = Pt(w)
    return sh


def chevron(sl, cx, cy, w=13, h=22):
    """o '>' fino entre os cartoes"""
    b = sl.shapes.build_freeform(Inches(I(cx - w / 2.0)), Inches(I(cy - h / 2.0)))
    b.add_line_segments([(Inches(I(cx + w / 2.0)), Inches(I(cy))),
                         (Inches(I(cx - w / 2.0)), Inches(I(cy + h / 2.0)))],
                        close=False)
    sh = stroke(b.convert_to_shape(), 1.5, CHEV)
    sh.name = 'Sequencia'
    return sh


# ------------------------------------------------------------------ conteudo
EYEBROW = 'L2S  |  LEAD QUENTE'
TITULO = [
    [('Converter o lead quente exige rodar quatro passos:', {})],
    [('comunicar, priorizar, avisar da consequência e realocar',
      {'color': RED, 'bold': True})],
]
SUBTITULO = [
    [('Pré-requisito: o número existe e é nominal: leads quentes recebidos, '
      'conversão e tempo até o primeiro contato', {})],
    [('por divisional, loja e CV', {})],
]
MERITO = ('Lead quente passa a ser mérito, não direito: quem converte recebe '
          'mais, quem não converte recebe menos')
FONTE = ('Fonte: análise Bain; discussão com liderança comercial '
         'Localiza Seminovos')

PASSOS = [
    {
        'n': '1',
        'icone': 'barras_linha',
        'titulo': ['Comunicar', 'divisionais'],
        'bullets': [
            'Glauco leva o número aos divisionais: leads quentes recebidos '
            'e conversão por regional',
            'Ranking nominal explícito: quem está bem e quem está mal',
            'Mensagem única: o acompanhamento passa a ser semanal',
        ],
        'dono': 'Dono: Glauco',
        'prazo': 'Prazo: Semana 1',
    },
    {
        'n': '2',
        'icone': 'pessoas_grupo_linha',
        'titulo': ['Priorizar', 'na ponta'],
        'bullets': [
            'Cascata completa: divisional para regional, regional para GV, '
            'GV para CV',
            'Cada CV recebe a lista nominal dos leads quentes do dia',
            'Instrução única: ligar em 100% da lista no mesmo dia',
        ],
        'dono': 'Dono: Regionais e GVs',
        'prazo': 'Prazo: Semanas 1 e 2',
    },
    {
        'n': '3',
        'icone': 'alerta_linha',
        'titulo': ['Avisar da', 'consequência'],
        'bullets': [
            'Regra comunicada antes de valer: não converter lead quente '
            'reduz o volume recebido',
            'Conversão por CV acompanhada semana a semana',
            'Sem surpresa: aviso formal antes do primeiro corte',
        ],
        'dono': 'Dono: Glauco e GVs',
        'prazo': 'Prazo: Semana 2',
    },
    {
        'n': '4',
        'icone': 'ciclo_linha',
        'titulo': ['Realocar', 'o lead quente'],
        'bullets': [
            'Alocação inicial antes o critério do cliente, como é hoje',
            'Dentro da loja, o mix de quente, morno e frio segue a '
            'performance do CV',
            'Lead quente não tratado no prazo é realocado para quem atende',
        ],
        'dono': 'Dono: L2S e comercial',
        'prazo': 'Prazo: Semanas 3 e 4, em piloto',
    },
]


# ------------------------------------------------------------------ pagina
def cabecalho(sl):
    txt(sl, MX, EYE_Y, 600, 16, EYEBROW, size=pt(12), color=GREY4, bold=True,
        spc=375, anchor=MSO_ANCHOR.TOP, name='Eyebrow')
    rect(sl, MX, BAR_Y, BAR_W, BAR_H, RED, None, name='Barra vermelha')
    txt(sl, MX, TITLE_Y, CW, TITLE_H, TITULO, size=pt(27), color=BLACK,
        lsp=1.05, name='Titulo')
    txt(sl, MX, SUB_Y, CW, SUB_H, SUBTITULO, size=pt(15), color=GREY2,
        lsp=1.18, name='Subtitulo')


def cartao(sl, i, passo):
    x = CARD_X[i]
    gradient_card(sl, x, CARD_Y, CARD_W, CARD_H, name='Passo %s' % passo['n'])
    icon(sl, passo['icone'], I(x + (CARD_W - ICON_D) / 2.0),
         I(CARD_Y + 52), I(ICON_D), RED)
    txt(sl, x + 20, CARD_Y + 114, CARD_W - 40, 44,
        [[(passo['titulo'][0], {})], [(passo['titulo'][1], {})]],
        size=pt(21), color=BLACK, bold=True, align=PP_ALIGN.CENTER, lsp=1.0,
        name='Passo %s titulo' % passo['n'])
    txt(sl, x + 20, CARD_Y + 162, CARD_W - 40, 162,
        [[(b, {})] for b in passo['bullets']],
        size=pt(14), color=BODY, lsp=1.12, space_after=pt(6), bullet=True,
        name='Passo %s bullets' % passo['n'])
    txt(sl, x + 20, CARD_Y + CARD_H - 44, CARD_W - 40, 34,
        [[(passo['dono'], {})], [(passo['prazo'], {})]],
        size=pt(14), color=BLACK, bold=True, lsp=1.18,
        anchor=MSO_ANCHOR.BOTTOM, name='Passo %s dono' % passo['n'])
    b = rect(sl, x + 6, CARD_Y - 5, BUBBLE_D, BUBBLE_D, RED, None,
             shape=MSO_SHAPE.OVAL, name='Numero %s' % passo['n'])
    label(b, passo['n'], size=pt(18), color=WHITE)


def faixa_merito(sl):
    rect(sl, MX, BAN_Y, CW, BAN_H, None, ROSE, 1.5, MSO_SHAPE.ROUNDED_RECTANGLE,
         adj=7.0 / BAN_H, name='Faixa de merito')
    rect(sl, MX - 8, BAN_Y - 7, MEDAL_D, MEDAL_D, RED, None,
         shape=MSO_SHAPE.OVAL, name='Medalha disco')
    icon(sl, 'trofeu_linha', I(MX - 8 + (MEDAL_D - 28) / 2.0),
         I(BAN_Y - 7 + (MEDAL_D - 28) / 2.0), I(28), WHITE)
    txt(sl, MX + 68, BAN_Y, CW - 88, BAN_H, MERITO, size=pt(16), color=RED,
        bold=True, anchor=MSO_ANCHOR.MIDDLE, name='Merito')


def rodape(sl):
    txt(sl, MX, FOOT_Y, 900, 16, FONTE, size=pt(10), color=GREY1,
        anchor=MSO_ANCHOR.MIDDLE, name='Fonte')
    txt(sl, MX + CW - 300, FOOT_Y, 238, 16, 'BAIN & COMPANY', size=pt(12),
        color=RED, bold=True, spc=200, align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE, name='Marca')
    rect(sl, MX + CW - 48, FOOT_Y + 1, 15, 15, None, RED, 1.5, MSO_SHAPE.OVAL,
         name='Marca simbolo')
    rect(sl, MX + CW - 37, FOOT_Y - 2, 4, 4, RED, None, shape=MSO_SHAPE.OVAL,
         name='Marca ponto')
    txt(sl, MX + CW - 14, FOOT_Y, 14, 16, '1', size=pt(11), color=GREY3,
        align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, name='Pagina')


def pagina(prs):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    cabecalho(sl)
    for i, passo in enumerate(PASSOS):
        cartao(sl, i, passo)
    for i in range(3):
        chevron(sl, CARD_X[i] + CARD_W + GAP / 2.0, CHEV_CY)
    faixa_merito(sl)
    rodape(sl)
    return sl


def main(dst):
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    pagina(prs)
    cp = prs.core_properties
    cp.title = 'Lead quente: comunicar, priorizar, avisar e realocar'
    cp.author = 'ConsultingInnovation@bain.com'
    cp.last_modified_by = 'ConsultingInnovation@bain.com'
    prs.save(dst)
    print('ok %s (%d pagina)' % (dst, len(prs.slides._sldIdLst)))


if __name__ == '__main__':
    import sys
    main(sys.argv[1] if len(sys.argv) > 1 else 'output/lead_quente.pptx')
