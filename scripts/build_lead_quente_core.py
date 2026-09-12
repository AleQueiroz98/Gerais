# -*- coding: utf-8 -*-
"""Pagina 'L2S | Lead quente' - os quatro passos para converter o lead quente.

Replica a pagina HTML `source/lead_quente.html` no template Bain Core (16:9),
trocando os elementos improvisados do HTML pelos elementos btfp oficiais:

  * circulos numerados  -> btfpNumberBubble
  * chevrons "›"        -> btfpSequenceArrow
  * faixa com medalha   -> btfpConclusionArrowHorizontal

Os quatro icones sao desenhados em line-art vermelho reproduzindo os SVGs do
HTML: nao sao icones da Toolbox CS, entao podem ser trocados no PowerPoint.

Requer os assets do skill bain-slides (template Bain Core + elementos btfp).
Aponte BAIN_SLIDES_DIR para a pasta do skill antes de rodar:

    BAIN_SLIDES_DIR=/caminho/para/bain-slides python3 scripts/build_lead_quente_core.py

Saida: output/lead_quente_core.pptx (uma pagina).
"""
import os
import sys
import io
import shutil
import zipfile
import pathlib

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = pathlib.Path(__file__).resolve().parent.parent
SK = pathlib.Path(os.environ.get("BAIN_SLIDES_DIR", "")).expanduser()
if not (SK / "templates" / "Bain_Core_On_Screen_Show__16_9_.pptx").exists():
    sys.exit("Defina BAIN_SLIDES_DIR apontando para a pasta do skill bain-slides.")

sys.path.insert(0, str(SK / "scripts"))
from load_bain_contract import load_bain_contract          # noqa: E402
from extract_master_contract import find_title_only_idx    # noqa: E402
from copy_elements import copy_shapes_batch                # noqa: E402
from dedupe_masters import dedupe_masters                  # noqa: E402

MASTER = str(SK / "templates" / "Bain_Core_On_Screen_Show__16_9_.pptx")
BTFP = str(SK / "templates" / "bain-btfp-elements") + os.sep
TMP = str(ROOT / "output" / ".lead_quente_core_tmp.pptx")
OUT = str(ROOT / "output" / "lead_quente_core.pptx")

contract = load_bain_contract()
PX = contract['px_to_emu']; SAFE = contract['safe']; TITLE = contract['title']
COLORS = contract['colors']; FONTS = contract['fonts']
RED = COLORS['primary']; BORDER = '#D6D6D6'
def px(n): return int(round(n*PX))
def rgb(h):
    h=h.lstrip('#'); return RGBColor(int(h[0:2],16),int(h[2:4],16),int(h[4:6],16))
A='http://schemas.openxmlformats.org/drawingml/2006/main'
P_NS='http://schemas.openxmlformats.org/presentationml/2006/main'

prs = Presentation(MASTER)
title_only = prs.slide_layouts[find_title_only_idx(prs)]
slide = prs.slides.add_slide(title_only)

def nobullet(p):
    pPr=p._p.get_or_add_pPr()
    for tag in (f'{{{A}}}buNone',f'{{{A}}}buChar',f'{{{A}}}buAutoNum'):
        el=pPr.find(tag)
        if el is not None: pPr.remove(el)
    etree.SubElement(pPr,f'{{{A}}}buNone')

def set_title(slide, text, size_pt=24):
    ph=slide.placeholders[0]; tf=ph.text_frame; tf.clear()
    p=tf.paragraphs[0]; nobullet(p)
    r=p.add_run(); r.text=text
    r.font.name=FONTS['title']; r.font.bold=False; r.font.size=Pt(size_pt)
    r.font.color.rgb=rgb(TITLE.get('color','#000000'))
    return ph

def add_text(slide,x,y,w,h,paras,align=None,anchor=None,spacing=None,line_sp=None):
    tb=slide.shapes.add_textbox(px(x),px(y),px(w),px(h))
    tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    if anchor: tf.vertical_anchor=anchor
    first=True
    for text,size,bold,color,sb in paras:
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        nobullet(p)
        pPr=p._p.get_or_add_pPr(); pPr.set('marL','0'); pPr.set('indent','0')
        p.space_before=Pt(sb or 0); p.space_after=Pt(0)
        if align: p.alignment=align
        if line_sp: p.line_spacing=line_sp
        r=p.add_run(); r.text=text
        r.font.size=Pt(size); r.font.bold=bold; r.font.name=FONTS['body']
        if color: r.font.color.rgb=rgb(color)
        if spacing:
            r.font._rPr.set('spc', str(int(spacing*100)))
    return tb

def add_bullets(slide,x,y,w,h,items,size=12,color='#1D1D1D',space=6,line_sp=0.92):
    tb=slide.shapes.add_textbox(px(x),px(y),px(w),px(h))
    tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    first=True
    for t in items:
        p=tf.paragraphs[0] if first else tf.add_paragraph()
        p.space_before=Pt(0 if first else space); p.space_after=Pt(0)
        first=False
        pPr=p._p.get_or_add_pPr()
        pPr.set('marL','133350'); pPr.set('indent','-133350')
        for tag in (f'{{{A}}}buNone',f'{{{A}}}buChar',f'{{{A}}}buAutoNum'):
            el=pPr.find(tag)
            if el is not None: pPr.remove(el)
        bf=etree.SubElement(pPr,f'{{{A}}}buFont'); bf.set('typeface','Arial'); bf.set('pitchFamily','34'); bf.set('charset','0')
        bc=etree.SubElement(pPr,f'{{{A}}}buChar'); bc.set('char','•')
        p.line_spacing=line_sp
        r=p.add_run(); r.text=t
        r.font.size=Pt(size); r.font.name=FONTS['body']; r.font.color.rgb=rgb(color)
    return tb

def shape(kind,x,y,w,h,fill=None,line=None,lw=2.0):
    s=slide.shapes.add_shape(kind,px(x),px(y),px(w),px(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb=rgb(fill)
    else: s.fill.background()
    if line:
        s.line.color.rgb=rgb(line); s.line.width=Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit=False
    if s.has_text_frame:
        s.text_frame.text=''
        nobullet(s.text_frame.paragraphs[0])
    return s

# ── layout constants ─────────────────────────────────────────────
CARD_W=[260,260,260,259]
CARD_X=[35,352,669,986]
CARD_Y=154; CARD_H=424
GAP_MID=[323.5,640.5,957.5]

# ── title block ──────────────────────────────────────────────────
add_text(slide,35,97,600,14,[("L2S | LEAD QUENTE",9,True,'#6B6B6B',0)],spacing=4)
set_title(slide,"Converter o lead quente exige rodar quatro passos: comunicar, priorizar, avisar da consequência e realocar")
add_text(slide,35,113,1180,18,
         [("Pré-requisito: o número existe e é nominal: leads quentes recebidos, conversão e tempo até o primeiro contato por divisional, loja e CV",11,False,'#474747',0)],
         line_sp=1.05)

# ── cards ────────────────────────────────────────────────────────
for i in range(4):
    shape(MSO_SHAPE.RECTANGLE,CARD_X[i],CARD_Y,CARD_W[i],CARD_H,fill='#FFFFFF',line=BORDER,lw=0.75)

HEADERS=[["Comunicar","divisionais"],["Priorizar","na ponta"],["Avisar da","consequência"],["Realocar","o lead quente"]]
BULLETS=[
 ["Glauco leva o número aos divisionais: leads quentes recebidos e conversão por regional",
  "Ranking nominal explícito: quem está bem e quem está mal",
  "Mensagem única: o acompanhamento passa a ser semanal"],
 ["Cascata completa: divisional para regional, regional para GV, GV para CV",
  "Cada CV recebe a lista nominal dos leads quentes do dia",
  "Instrução única: ligar em 100% da lista no mesmo dia"],
 ["Regra comunicada antes de valer: não converter lead quente reduz o volume recebido",
  "Conversão por CV acompanhada semana a semana",
  "Sem surpresa: aviso formal antes do primeiro corte"],
 ["Alocação inicial antes o critério do cliente, como é hoje",
  "Dentro da loja, o mix de quente, morno e frio segue a performance do CV",
  "Lead quente não tratado no prazo é realocado para quem atende"]]
OWNERS=[["Dono: Glauco","Prazo: Semana 1"],["Dono: Regionais e GVs","Prazo: Semanas 1 e 2"],
        ["Dono: Glauco e GVs","Prazo: Semana 2"],["Dono: L2S e comercial","Prazo: Semanas 3 e 4, em piloto"]]

# ── icons (red line art, drawn) ──────────────────────────────────
def icon_bars(bx,by):
    for dx,dy,h in ((3,34,22),(22,18,38),(41,4,52)):
        shape(MSO_SHAPE.RECTANGLE,bx+dx,by+dy,12,h,fill=None,line=RED,lw=2.0)
def icon_people(bx,by):
    shape(MSO_SHAPE.OVAL,bx+21,by+2,16,16,fill=None,line=RED,lw=2.0)
    shape(MSO_SHAPE.OVAL,bx+2,by+12,13,13,fill=None,line=RED,lw=2.0)
    shape(MSO_SHAPE.OVAL,bx+43,by+12,13,13,fill=None,line=RED,lw=2.0)
    shape(MSO_SHAPE.BLOCK_ARC,bx+14,by+22,30,34,fill=None,line=RED,lw=2.0)
    shape(MSO_SHAPE.BLOCK_ARC,bx-3,by+30,24,28,fill=None,line=RED,lw=2.0)
    shape(MSO_SHAPE.BLOCK_ARC,bx+37,by+30,24,28,fill=None,line=RED,lw=2.0)
def icon_warn(bx,by):
    shape(MSO_SHAPE.ISOSCELES_TRIANGLE,bx+1,by+4,56,50,fill=None,line=RED,lw=2.0)
    shape(MSO_SHAPE.RECTANGLE,bx+27,by+22,3,17,fill=RED,line=None)
    shape(MSO_SHAPE.OVAL,bx+26.5,by+43,4,4,fill=RED,line=None)
def icon_cycle(bx,by):
    a=shape(MSO_SHAPE.CIRCULAR_ARROW,bx+2,by+2,54,54,fill=None,line=RED,lw=2.0)
    b=shape(MSO_SHAPE.CIRCULAR_ARROW,bx+2,by+2,54,54,fill=None,line=RED,lw=2.0)
    b.rotation=180
ICONS=[icon_bars,icon_people,icon_warn,icon_cycle]

for i in range(4):
    cx=CARD_X[i]; cw=CARD_W[i]
    ICONS[i](cx+cw/2-29, CARD_Y+24)
    add_text(slide,cx+18,CARD_Y+92,cw-36,44,
             [(HEADERS[i][0],16,True,'#000000',0),(HEADERS[i][1],16,True,'#000000',0)],
             align=PP_ALIGN.CENTER,line_sp=0.95)
    add_bullets(slide,cx+16,CARD_Y+142,cw-32,214,BULLETS[i])
    add_text(slide,cx+18,CARD_Y+368,cw-36,48,
             [(OWNERS[i][0],11,True,'#000000',0),(OWNERS[i][1],11,True,'#000000',0)],line_sp=1.0)

# ── source ───────────────────────────────────────────────────────
add_text(slide,35,586,900,14,
         [("Fonte: análise Bain; discussão com liderança comercial Localiza Seminovos",8,False,'#000000',0)])

prs.save(TMP)

# ------------------------------------------ elementos btfp oficiais
DECK = TMP
TGT=2


specs=[]
for i in range(4):
    specs.append({'source':BTFP+'btfpNumberBubble.pptx','slide':1,'shape':'btfpNumberBubble469252','target_slide':TGT})
for i in range(3):
    specs.append({'source':BTFP+'btfpSequenceArrow.pptx','slide':1,'shape':'btfpSequenceArrow864161','target_slide':TGT})
specs.append({'source':BTFP+'btfpConclusionArrowHorizontal.pptx','slide':1,'shape':'btfpConclusionArrow779861','target_slide':TGT})

copy_shapes_batch(None, specs, DECK)

prs=Presentation(DECK)
slide=prs.slides[TGT]
shapes=list(slide.shapes)
added=shapes[-8:]
bubbles=added[0:4]; arrows=added[4:7]; concl=added[7]

for i,b in enumerate(bubbles):
    b.left=px(CARD_X[i]+8); b.top=px(CARD_Y-16); b.width=px(33); b.height=px(33)
    for t in b._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'):
        t.text=str(i+1)
    for rPr in b._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr'):
        rPr.set('sz','1600')

for i,a in enumerate(arrows):
    a.width=px(18); a.height=px(70)
    a.left=px(GAP_MID[i]-9); a.top=px(170)

# conclusion arrow: keep native geometry (x=35, y=603, w=1211, h=87)
TXT="Lead quente passa a ser mérito, não direito: quem converte recebe mais, quem não converte recebe menos"
for t in concl._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'):
    if t.text and 'conclusion' in t.text.lower():
        t.text=TXT

prs.save(DECK)

# ---------------------------------------------------------------- 1 pagina
dedupe_masters(TMP)
shutil.copy(TMP, OUT)
with zipfile.ZipFile(OUT, 'r') as z:
    files = {n: z.read(n) for n in z.namelist()}
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
root = etree.fromstring(files['ppt/presentation.xml'])
lst = root.find('.//p:sldIdLst', NS)
for el in list(lst)[:-1]:
    lst.remove(el)
files['ppt/presentation.xml'] = etree.tostring(
    root, xml_declaration=True, encoding='UTF-8', standalone=True)
rels = etree.fromstring(files['ppt/_rels/presentation.xml.rels'])
slide_rels = [e for e in list(rels)
              if 'slide' in e.get('Target', '')
              and 'slideLayout' not in e.get('Target', '')
              and 'slideMaster' not in e.get('Target', '')]
for e in slide_rels[:-1]:
    rels.remove(e)
files['ppt/_rels/presentation.xml.rels'] = etree.tostring(
    rels, xml_declaration=True, encoding='UTF-8', standalone=True)
buf = io.BytesIO()
with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zo:
    for name, data in files.items():
        zo.writestr(name, data)
open(OUT, 'wb').write(buf.getvalue())
os.remove(TMP)
print("Gerado:", OUT, "-", len(Presentation(OUT).slides), "pagina")
