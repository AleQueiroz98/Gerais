# -*- coding: utf-8 -*-
"""Cinco paginas 16:9 no padrao Bain Core, reproduzindo os prints enviados.

  1. LEAD QUENTE | NOSSA PROPOSTA  - quatro passos em cards, dono e prazo
  2. F&I MULTIBANCO | VISAO GERAL  - ambicao, valor para e cinco frentes
  3. F&I MULTIBANCO | PLANO DE EXECUCAO - tabela das cinco frentes
  4. L2S | LEAD QUENTE             - variante da pagina 1, com icones e logo
  5. F&I MULTIBANCO                - variante da pagina 2, com icones e logo

Tudo e desenhado com shapes, caixas de texto e tabela nativas do PowerPoint:
nada de imagem colada, todo elemento continua editavel no ppt.

    python3 build_lead_quente_fi.py ../output/lead_quente_fi_multibanco.pptx
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from PIL import ImageFont

# ------------------------------------------------------------------ tokens
RED    = RGBColor(0xCC, 0x00, 0x00)   # Bain red
RED_D  = RGBColor(0x99, 0x00, 0x00)
RED_BG = RGBColor(0xFA, 0xEC, 0xEC)
BLACK  = RGBColor(0x00, 0x00, 0x00)
GREY1  = RGBColor(0x33, 0x33, 0x33)
GREY2  = RGBColor(0x5C, 0x5C, 0x5C)
GREY3  = RGBColor(0x85, 0x85, 0x85)
GREY4  = RGBColor(0xB4, 0xB4, 0xB4)
GREY5  = RGBColor(0xD6, 0xD6, 0xD6)
CARD   = RGBColor(0xF2, 0xF2, 0xF2)
BAND   = RGBColor(0xE8, 0xE8, 0xE8)
BAND2  = RGBColor(0xED, 0xED, 0xED)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

FONT   = 'Arial'
LANG   = 'pt-BR'
SW, SH = 13.333, 7.5
M      = 0.55                 # margem esquerda/direita
CW     = SW - 2 * M           # largura util

T_EYE  = 8.0                  # kicker
T_TIT  = 19.0                 # titulo de acao
T_SUB  = 10.5
T_FOOT = 7.5

_TTF = {False: '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
        True:  '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf'}
_fc = {}


def text_w(s, pt, bold=False):
    """largura do texto em polegadas"""
    if bold not in _fc:
        _fc[bold] = ImageFont.truetype(_TTF[bold], 200)
    return _fc[bold].getlength(s) / 200.0 * pt / 72.0


# ------------------------------------------------------------------ helpers
def no_shadow(shape):
    el = shape._element
    st = el.find(qn('p:style'))
    if st is not None:
        el.remove(st)
    spPr = el.spPr
    for e in spPr.findall(qn('a:effectLst')):
        spPr.remove(e)
    spPr.append(spPr.makeelement(qn('a:effectLst'), {}))


def shp(sl, kind, x, y, w, h, fill=None, line=None, lw=1.0, rot=0.0, adj=None):
    s = sl.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    no_shadow(s)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    if rot:
        s.rotation = rot
    if adj is not None:
        for i, v in enumerate(adj):
            try:
                s.adjustments[i] = v
            except (IndexError, ValueError):
                pass
    s.text_frame.word_wrap = True
    return s


def rect(sl, x, y, w, h, fill=None, line=None, lw=1.0, rot=0.0):
    return shp(sl, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, lw, rot)


def ell(sl, x, y, w, h, fill=None, line=None, lw=1.0):
    return shp(sl, MSO_SHAPE.OVAL, x, y, w, h, fill, line, lw)


def hline(sl, x, y, w, color=GREY5, thick=0.012):
    return rect(sl, x, y, w, thick, color, None)


def vline(sl, x, y, h, color=GREY5, thick=0.010):
    return rect(sl, x, y, thick, h, color, None)


def _run(p, text, size, color, bold=False, italic=False, spc=None, font=FONT):
    r = p.add_run()
    r.text = text
    f = r.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    rPr.set('lang', LANG)
    rPr.set('dirty', '0')
    if spc:
        rPr.set('spc', str(int(spc)))
    return r


def _norm(lines):
    """str | [(txt, bold)] | [[(txt, bold)], ...]  ->  lista de paragrafos"""
    if isinstance(lines, str):
        return [[(lines, False)]]
    if lines and isinstance(lines[0], tuple):
        return [list(lines)]
    return [list(p) if not isinstance(p, str) else [(p, False)] for p in lines]


def txt(host, x=None, y=None, w=None, h=None, lines='', size=10, color=GREY1,
        bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, lsp=1.15,
        spc=None, space_before=0.0, bullet=False, indent=0.14, wrap=True):
    """host: slide (cria textbox) ou shape (usa o text_frame existente)"""
    if x is None:
        tf = host.text_frame
        tf.word_wrap = True
    else:
        box = host.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = _norm(lines)
    for i, parts in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = lsp
        if i and space_before:
            p.space_before = Pt(space_before * 72)
        for seg, b in parts:
            _run(p, seg, size, color, bold or b, spc=spc)
        if bullet:
            _bullet(p, indent)
    return tf


def _bullet(p, indent=0.14):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(indent * 914400)))
    pPr.set('indent', str(int(-indent * 914400)))
    bf = pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'})
    bc = pPr.makeelement(qn('a:buChar'), {'char': '•'})
    pPr.append(bf)
    pPr.append(bc)


def bullets(sl, x, y, w, items, size=10, color=GREY1, gap=0.13, lsp=1.16,
            indent=0.14):
    tf = txt(sl, x, y, w, 0.4, [[(i, False)] if isinstance(i, str) else i
                                for i in items],
             size=size, color=color, lsp=lsp, space_before=gap)
    for p in tf.paragraphs:
        _bullet(p, indent)
    return tf


# ------------------------------------------------------------------ chrome
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def eyebrow(sl, text, rule=False):
    txt(sl, M, 0.26, CW, 0.2, [[(text.upper(), False)]], size=T_EYE,
        color=GREY3, bold=False, spc=170)
    if rule:
        hline(sl, M, 0.50, CW, GREY5, 0.010)


def title(sl, l1, l2, y=0.56, size=T_TIT):
    txt(sl, M, y, CW, 0.40, [[(l1, True)]], size=size, color=BLACK, bold=True,
        lsp=1.0)
    txt(sl, M, y + 0.345, CW, 0.40, [[(l2, True)]], size=size, color=RED,
        bold=True, lsp=1.0)


def footer(sl, source, page, logo=False):
    y = 7.02
    hline(sl, M, y - 0.10, CW, GREY5, 0.008)
    txt(sl, M, y, 10.6, 0.3, [[('Fonte: ', True), (source, False)]],
        size=T_FOOT, color=GREY3)
    if logo:
        r = SW - M
        txt(sl, r - 0.34, y, 0.34, 0.3, [[('%d' % page, False)]],
            size=T_FOOT, color=GREY3, align=PP_ALIGN.RIGHT, wrap=False)
        txt(sl, r - 0.58, y, 0.14, 0.3, [[('|', False)]], size=T_FOOT,
            color=GREY4, align=PP_ALIGN.CENTER, wrap=False)
        ell(sl, r - 0.86, y + 0.005, 0.145, 0.145, RED, None)
        lw = text_w('BAIN & COMPANY', 8.0, True) + 0.22
        txt(sl, r - 0.96 - lw, y - 0.005, lw, 0.3,
            [[('BAIN & COMPANY', True)]], size=8.0, color=BLACK, bold=True,
            spc=60, align=PP_ALIGN.RIGHT, wrap=False)
    else:
        txt(sl, SW - M - 0.6, y, 0.6, 0.3, [[('%d' % page, False)]],
            size=T_FOOT, color=GREY3, align=PP_ALIGN.RIGHT)


def numbubble(sl, cx, cy, d, n, fill=RED, fg=WHITE, size=None):
    c = ell(sl, cx, cy, d, d, fill, None)
    txt(c, lines=[[(str(n), True)]], size=size or d * 36, color=fg, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, lsp=1.0)
    return c


# ------------------------------------------------------------------ icones
def ic_bars(sl, x, y, s, c):
    for dx, hf in ((0.00, 0.42), (0.33, 0.72), (0.66, 1.0)):
        rect(sl, x + dx * s, y + s * (1 - hf), s * 0.24, s * hf, c, None)


def ic_person(sl, x, y, s, c):
    ell(sl, x + s * 0.30, y + s * 0.02, s * 0.40, s * 0.40, None, c, 1.5)
    b = shp(sl, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x + s * 0.10, y + s * 0.52,
            s * 0.80, s * 0.46, None, c, 1.5, adj=(0.45,))
    return b


def ic_people(sl, x, y, s, c):
    for dx, dy, sc in ((0.00, 0.06, 0.86), (0.33, 0.00, 1.0),
                       (0.66, 0.06, 0.86)):
        bw, bh = s * 0.34 * sc, s * 0.40 * sc
        bx, by = x + dx * s, y + (0.48 + dy) * s
        d = s * 0.27 * sc
        ell(sl, bx + (bw - d) / 2, by - d + s * 0.03, d, d, None, c, 1.5)
        shp(sl, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, bx, by, bw, bh,
            None, c, 1.5, adj=(0.48,))


def ic_warn(sl, x, y, s, c):
    shp(sl, MSO_SHAPE.ISOSCELES_TRIANGLE, x, y + s * 0.06, s, s * 0.88,
        None, c, 1.6, adj=(0.5,))
    rect(sl, x + s * 0.47, y + s * 0.38, s * 0.06, s * 0.28, c, None)
    rect(sl, x + s * 0.47, y + s * 0.72, s * 0.06, s * 0.07, c, None)


def ic_refresh(sl, x, y, s, c):
    shp(sl, MSO_SHAPE.CIRCULAR_ARROW, x, y, s, s, c, None,
        adj=(0.10, 4.7e6, 1.2e6, 1.05e7, 0.055))


def ic_bank(sl, x, y, s, c):
    shp(sl, MSO_SHAPE.ISOSCELES_TRIANGLE, x + s * 0.02, y + s * 0.04,
        s * 0.96, s * 0.26, c, None, adj=(0.5,))
    for i in range(4):
        rect(sl, x + s * (0.12 + i * 0.24), y + s * 0.36, s * 0.10, s * 0.42,
             c, None)
    rect(sl, x, y + s * 0.84, s, s * 0.13, c, None)


def ic_gear(sl, x, y, s, c):
    shp(sl, MSO_SHAPE.GEAR_9, x, y, s, s, c, None)
    ell(sl, x + s * 0.33, y + s * 0.33, s * 0.34, s * 0.34, WHITE, None)


def ic_monitor(sl, x, y, s, c):
    rect(sl, x, y + s * 0.06, s, s * 0.62, None, c, 1.6)
    rect(sl, x + s * 0.12, y + s * 0.20, s * 0.48, s * 0.06, c, None)
    rect(sl, x + s * 0.12, y + s * 0.34, s * 0.72, s * 0.06, c, None)
    rect(sl, x + s * 0.12, y + s * 0.48, s * 0.36, s * 0.06, c, None)
    rect(sl, x + s * 0.44, y + s * 0.68, s * 0.12, s * 0.18, c, None)
    rect(sl, x + s * 0.22, y + s * 0.86, s * 0.56, s * 0.09, c, None)


def ic_share(sl, x, y, s, c):
    rect(sl, x + s * 0.18, y + s * 0.24, s * 0.62, s * 0.045, c, None, rot=-32)
    rect(sl, x + s * 0.18, y + s * 0.70, s * 0.62, s * 0.045, c, None, rot=32)
    d = s * 0.28
    ell(sl, x + s * 0.72, y, d, d, c, None)
    ell(sl, x, y + s * 0.36, d, d, c, None)
    ell(sl, x + s * 0.72, y + s * 0.72, d, d, c, None)


def ic_target(sl, x, y, s, c):
    shp(sl, MSO_SHAPE.DONUT, x, y, s, s, c, None, adj=(0.12,))
    shp(sl, MSO_SHAPE.DONUT, x + s * 0.26, y + s * 0.26, s * 0.48, s * 0.48,
        c, None, adj=(0.22,))
    ell(sl, x + s * 0.42, y + s * 0.42, s * 0.16, s * 0.16, c, None)


def ic_trophy(sl, x, y, s, c):
    shp(sl, MSO_SHAPE.DONUT, x + s * 0.02, y + s * 0.06, s * 0.26, s * 0.30,
        c, None, adj=(0.22,))
    shp(sl, MSO_SHAPE.DONUT, x + s * 0.72, y + s * 0.06, s * 0.26, s * 0.30,
        c, None, adj=(0.22,))
    shp(sl, MSO_SHAPE.TRAPEZOID, x + s * 0.24, y + s * 0.02, s * 0.52,
        s * 0.48, c, None, rot=180, adj=(0.20,))
    rect(sl, x + s * 0.45, y + s * 0.50, s * 0.10, s * 0.22, c, None)
    rect(sl, x + s * 0.28, y + s * 0.72, s * 0.44, s * 0.09, c, None)
    rect(sl, x + s * 0.20, y + s * 0.83, s * 0.60, s * 0.11, c, None)


def ic_car(sl, x, y, s, c):
    shp(sl, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x + s * 0.22, y + s * 0.14,
        s * 0.56, s * 0.32, c, None, adj=(0.42,))
    shp(sl, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x, y + s * 0.44, s, s * 0.26,
        c, None, adj=(0.34,))
    ell(sl, x + s * 0.13, y + s * 0.62, s * 0.24, s * 0.24, c, None)
    ell(sl, x + s * 0.63, y + s * 0.62, s * 0.24, s * 0.24, c, None)
    ell(sl, x + s * 0.20, y + s * 0.69, s * 0.10, s * 0.10, WHITE, None)
    ell(sl, x + s * 0.70, y + s * 0.69, s * 0.10, s * 0.10, WHITE, None)


def ic_db(sl, x, y, s, c):
    for i in range(3):
        shp(sl, MSO_SHAPE.CAN, x + s * 0.10, y + s * (0.02 + i * 0.32),
            s * 0.80, s * 0.32, c, None, adj=(0.30,))


ICONS = {'bars': ic_bars, 'person': ic_person, 'people': ic_people,
         'warn': ic_warn, 'refresh': ic_refresh, 'bank': ic_bank,
         'gear': ic_gear, 'monitor': ic_monitor, 'share': ic_share,
         'target': ic_target, 'trophy': ic_trophy, 'car': ic_car,
         'db': ic_db}


def icon(sl, kind, x, y, s, c=RED):
    ICONS[kind](sl, x, y, s, c)


# ------------------------------------------------------------------ conteudo
PASSOS = [
    dict(n=1, tit='Comunicar\ndivisionais', ic='bars', dono='Glauco',
         prazo='Semana 1', dono_ic='person', bullets=[
             [('Glauco leva o número aos divisionais: ', False),
              ('leads quentes recebidos e conversão por regional', False)],
             [('Ranking nominal explícito: ', False),
              ('quem está bem e quem está mal', False)],
             [('Mensagem única: ', False),
              ('o acompanhamento passa a ser semanal', False)]]),
    dict(n=2, tit='Priorizar\nna ponta', ic='people', dono='Regionais e GVs',
         prazo='Semanas 1 e 2', dono_ic='people', bullets=[
             [('Cascata completa: ', False),
              ('divisional para regional, regional para GV, GV para CV',
               False)],
             [('Cada CV recebe a lista nominal dos leads quentes do dia',
               False)],
             [('Instrução única: ', False),
              ('ligar em 100% da lista no mesmo dia', False)]]),
    dict(n=3, tit='Avisar da\nconsequência', ic='warn', dono='Glauco e GVs',
         prazo='Semana 2', dono_ic='person', bullets=[
             [('Regra comunicada antes de valer: ', False),
              ('não converter lead quente reduz o volume recebido', False)],
             [('Conversão por CV acompanhada semana a semana', False)],
             [('Sem surpresa: ', False),
              ('aviso formal antes do primeiro corte', False)]]),
    dict(n=4, tit='Realocar\no lead quente', ic='refresh',
         dono='L2S e comercial', prazo='Semanas 3 e 4, em piloto',
         dono_ic='people', bullets=[
             [('Alocação à loja segue o critério do cliente, como hoje',
               False)],
             [('Dentro da loja, o mix de quente, morno e frio segue a '
               'performance do CV', False)],
             [('Lead quente não tratado no prazo é realocado para quem '
               'atende', False)]]),
]

FRENTES = [
    dict(n=1, tit='Bancos', ic='bank',
         txt='Integrar em ondas, dos bancos menores aos maiores, à medida '
             'que a proposta de valor se confirma'),
    dict(n=2, tit='Tecnologia', ic='gear',
         txt='APIs para simulação via Corban do lojista e envio da ficha a '
             'múltiplos bancos em um único disparo'),
    dict(n=3, tit='Interface\nda plataforma', ic='monitor',
         txt='Melhoria contínua guiada pelo uso do vendedor, com ciclos '
             'curtos de release'),
    dict(n=4, tit='Lojistas', ic='people',
         txt='Piloto em pequena escala para provar a proposta de valor e, '
             'confirmada, escalar por onda'),
    dict(n=5, tit='Ecossistema', ic='share',
         txt='Novos serviços e integração com Valoriza depois que o núcleo '
             'estiver estável'),
]

VALOR = [
    dict(tit='Lojista e seu cliente', ic='people',
         txt='Mais aprovação, menos retrabalho e resposta em minutos'),
    dict(tit='Localiza', ic='car',
         txt='Mais penetração de F&I e relação mais forte com a rede de '
             'lojistas'),
    dict(tit='Banco', ic='bank',
         txt='Originação qualificada com custo de aquisição menor'),
]

AMBICAO = ('Plataforma de F&I multibanco e ponta a ponta que faz o lojista '
           'vender mais, com simulação e envio de ficha em um só fluxo')

PLANO = [
    ('1', 'Bancos', 'Bancos integrados\nem cascata',
     '1 banco conectado,\nenvio manual', '4 a 5 bancos em\ncascata automática',
     '2T26'),
    ('2', 'Tecnologia', 'APIs de simulação\ne envio de ficha',
     'Simulação via Corban,\nficha manual',
     'Disparo único para\nmúltiplos bancos', '2T26'),
    ('3', 'Interface da\nplataforma', 'Jornada do vendedor\nponta a ponta',
     'Versão inicial em uso', 'Release contínuo\nguiado pelo uso',
     'Contínuo'),
    ('4', 'Lojistas', 'Piloto e plano\nde escala',
     'Amostra do piloto\nem definição', 'Piloto provado e\nescala por onda',
     '1T26'),
    ('5', 'Ecossistema', 'Integração com\nValoriza e novos serviços',
     'Escopo em discussão', 'Oferta ampliada na\nmesma jornada', '2027'),
]

SRC_LQ = 'análise Bain; discussão com liderança comercial Localiza Seminovos'
SRC_F1 = 'análise Bain. Preliminar, para discussão'
SRC_F2 = ('análise Bain. Conteúdo preliminar, para preenchimento com os '
          'donos de cada frente')

TIT_LQ = ('Converter o lead quente exige rodar quatro passos:',
          'comunicar, priorizar, avisar da consequência e realocar')
TIT_FI = ('Entregar a plataforma de F&I multibanco depende de cinco frentes,',
          'iniciadas por um piloto com lojistas')

MERITO = ('Lead quente passa a ser mérito, não direito:',
          ' quem converte recebe mais, quem não converte recebe menos')


# ------------------------------------------------------------------ pagina 1
def page_lead_quente(prs):
    sl = blank(prs)
    eyebrow(sl, 'Lead quente  |  Nossa proposta')
    title(sl, *TIT_LQ)
    txt(sl, M, 1.30, 11.3, 0.5,
        [[('Pré-requisito: o número existe e é nominal – leads quentes '
           'recebidos, conversão e tempo até o primeiro contato por '
           'divisional, loja e CV', False)]],
        size=T_SUB, color=GREY2, lsp=1.20)

    cy, ch, gap = 2.02, 4.28, 0.22
    cwid = (CW - 3 * gap) / 4.0
    for i, p in enumerate(PASSOS):
        x = M + i * (cwid + gap)
        rect(sl, x, cy, cwid, ch, CARD, None)
        numbubble(sl, x + 0.20, cy + 0.22, 0.42, p['n'], size=14.5)
        txt(sl, x + 0.76, cy + 0.22, cwid - 0.96, 0.7,
            [[(l, True)] for l in p['tit'].split('\n')],
            size=13.5, color=BLACK, bold=True, lsp=1.10)
        bullets(sl, x + 0.22, cy + 1.10, cwid - 0.44, p['bullets'],
                size=9.5, color=GREY1, gap=0.13)
        # dono e prazo
        by = cy + ch - 0.82
        icon(sl, p['dono_ic'], x + 0.22, by + 0.02, 0.34)
        txt(sl, x + 0.72, by, cwid - 0.92, 0.6,
            [[('Dono: ', True), (p['dono'], False)],
             [('Prazo: ', True), (p['prazo'], False)]],
            size=9.5, color=GREY1, lsp=1.25, space_before=0.05)

    # faixa de merito
    fy = 6.50
    rect(sl, M, fy, CW, 0.46, CARD, None)
    c = ell(sl, M + 0.10, fy + 0.08, 0.30, 0.30, RED, None)
    txt(c, lines=[[('›', True)]], size=15, color=WHITE, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, lsp=1.0)
    txt(sl, M + 0.56, fy, CW - 0.70, 0.46,
        [[(MERITO[0], True), (MERITO[1], False)]],
        size=10.5, color=BLACK, anchor=MSO_ANCHOR.MIDDLE)

    footer(sl, SRC_LQ, 1)
    return sl


# ------------------------------------------------------------------ pagina 2
def page_fi_visao(prs):
    sl = blank(prs)
    eyebrow(sl, 'F&I multibanco  |  Visão geral')
    title(sl, *TIT_FI)

    # ambicao
    ay, ah = 1.40, 0.80
    lab = rect(sl, M, ay, 1.95, ah, RED, None)
    txt(lab, lines=[[('Ambição', True)]], size=12, color=WHITE, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    b = rect(sl, M + 2.05, ay, CW - 2.05, ah, BAND2, None)
    txt(sl, M + 2.30, ay, CW - 2.60, ah, [[(AMBICAO, True)]], size=11.5,
        color=GREY1, bold=True, anchor=MSO_ANCHOR.MIDDLE, lsp=1.22)

    # valor para
    vy, vh = 2.36, 1.34
    lab = rect(sl, M, vy, 1.95, vh, BAND, None)
    txt(lab, lines=[[('Valor para...', True)]], size=12, color=GREY1,
        bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    vx0, vgap = M + 2.05, 0.18
    vw = (CW - 2.05 - 2 * vgap) / 3.0
    for i, v in enumerate(VALOR):
        x = vx0 + i * (vw + vgap)
        rect(sl, x, vy, vw, vh, CARD, None)
        icon(sl, v['ic'], x + 0.22, vy + 0.20, 0.34)
        txt(sl, x + 0.70, vy + 0.18, vw - 0.90, 0.32, [[(v['tit'], True)]],
            size=11.5, color=BLACK, bold=True)
        txt(sl, x + 0.70, vy + 0.60, vw - 0.90, 0.66, [[(v['txt'], False)]],
            size=10, color=GREY1, lsp=1.20)

    # divisor
    dy = 4.02
    hline(sl, M, dy, CW, RED, 0.020)
    shp(sl, MSO_SHAPE.LEFT_ARROW, M - 0.06, dy - 0.075, 0.52, 0.17, RED, None,
        adj=(0.50, 0.42))
    lbl = 'cinco frentes, escaladas em ondas'
    lwd = text_w(lbl, 11.5, True) + 0.24
    rect(sl, M + 0.62, dy - 0.13, lwd, 0.26, WHITE, None)
    txt(sl, M + 0.74, dy - 0.115, lwd, 0.24, [[(lbl, True)]], size=11.5,
        color=RED, bold=True)

    # cinco frentes
    fy = 4.34
    fw = CW / 5.0
    for i, f in enumerate(FRENTES):
        x = M + i * fw
        if i:
            vline(sl, x - 0.005, fy + 0.06, 2.30, GREY5)
        numbubble(sl, x + 0.14, fy, 0.36, f['n'], size=12.5)
        icon(sl, f['ic'], x + 0.16, fy + 0.56, 0.34)
        txt(sl, x + 0.14, fy + 1.02, fw - 0.36, 0.6,
            [[(l, True)] for l in f['tit'].split('\n')],
            size=11.5, color=BLACK, bold=True, lsp=1.12)
        ty = fy + 1.52
        txt(sl, x + 0.14, ty, fw - 0.36, 1.1, [[(f['txt'], False)]],
            size=9.5, color=GREY1, lsp=1.20)

    footer(sl, SRC_F1, 2)
    return sl


# ------------------------------------------------------------------ pagina 3
def page_fi_plano(prs):
    sl = blank(prs)
    eyebrow(sl, 'F&I multibanco  |  Plano de execução')
    title(sl, 'As cinco frentes têm entregável, ponto de partida, destino',
          'e prazo explícitos')

    cols = [0.62, 1.78, 2.62, 2.62, 3.08, 1.51]
    heads = ['#', 'Frente', 'Entregável', 'Onde estamos',
             'Onde queremos chegar', 'Quando']
    ty, hh, rh = 1.62, 0.52, 0.88
    gf = sl.shapes.add_table(len(PLANO) + 1, len(cols), Inches(M), Inches(ty),
                             Inches(sum(cols)), Inches(hh + rh * len(PLANO)))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr
    for e in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(e)
    sid = tblPr.makeelement(qn('a:tableStyleId'), {})
    sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'   # No Style, No Grid
    tblPr.append(sid)
    for i, w in enumerate(cols):
        tbl.columns[i].width = Emu(int(round(w * 914400)))
    tbl.rows[0].height = Emu(int(round(hh * 914400)))
    for r in range(1, len(PLANO) + 1):
        tbl.rows[r].height = Emu(int(round(rh * 914400)))

    def cell(r, c, parts, size, color, bold=False, align=PP_ALIGN.LEFT,
             fill=None):
        cl = tbl.cell(r, c)
        cl.margin_left = cl.margin_right = Inches(0.14)
        cl.margin_top = cl.margin_bottom = Inches(0.06)
        cl.vertical_anchor = MSO_ANCHOR.MIDDLE
        if fill is None:
            cl.fill.background()
        else:
            cl.fill.solid()
            cl.fill.fore_color.rgb = fill
        tf = cl.text_frame
        tf.word_wrap = True
        tf.clear()
        for i, line in enumerate(parts.split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = 1.16
            _run(p, line, size, color, bold)
        return cl

    for c, h in enumerate(heads):
        cell(0, c, h, 10.5, GREY1, True,
             PP_ALIGN.CENTER if c in (0, 5) else PP_ALIGN.LEFT, BAND)
    for r, row in enumerate(PLANO, start=1):
        stripe = WHITE if r % 2 else RGBColor(0xFA, 0xFA, 0xFA)
        for c, v in enumerate(row):
            if c == 0:
                cell(r, c, '', 10, GREY1, False, PP_ALIGN.CENTER, stripe)
            elif c == 5:
                cell(r, c, v, 10, GREY1, False, PP_ALIGN.CENTER,
                     RGBColor(0xF4, 0xF4, 0xF4))
            else:
                cell(r, c, v, 10, BLACK if c == 1 else GREY1, c == 1,
                     PP_ALIGN.LEFT, stripe)
        hline(sl, M, ty + hh + rh * (r - 1), sum(cols), GREY5, 0.008)
        numbubble(sl, M + cols[0] / 2 - 0.18,
                  ty + hh + rh * (r - 1) + rh / 2 - 0.18, 0.36, row[0],
                  size=12.5)
    hline(sl, M, ty + hh + rh * len(PLANO), sum(cols), GREY5, 0.008)

    footer(sl, SRC_F2, 3)
    return sl


# ------------------------------------------------------------------ pagina 4
def page_lead_quente_v2(prs):
    sl = blank(prs)
    eyebrow(sl, 'L2S  |  Lead quente')
    title(sl, *TIT_LQ)
    txt(sl, M, 1.30, 11.3, 0.5,
        [[('Pré-requisito: o número existe e é nominal: leads quentes '
           'recebidos, conversão e tempo até o primeiro contato por '
           'divisional, loja e CV', False)]],
        size=T_SUB, color=GREY2, lsp=1.20)

    cy, ch, gap = 2.00, 4.28, 0.34
    cwid = (CW - 3 * gap) / 4.0
    hb = 1.02                       # altura da faixa de cabecalho do card
    for i, p in enumerate(PASSOS):
        x = M + i * (cwid + gap)
        rect(sl, x, cy + 0.20, cwid, hb, BAND2, None)
        icon(sl, p['ic'], x + cwid / 2 - 0.19, cy + 0.44, 0.38)
        numbubble(sl, x - 0.09, cy, 0.42, p['n'], size=14.5)
        txt(sl, x + 0.10, cy + hb + 0.30, cwid - 0.20, 0.62,
            [[(l, True)] for l in p['tit'].split('\n')],
            size=13.5, color=BLACK, bold=True, lsp=1.10)
        bullets(sl, x + 0.10, cy + hb + 0.98, cwid - 0.20, p['bullets'],
                size=9.5, color=GREY1, gap=0.13)
        by = cy + ch - 0.62
        txt(sl, x + 0.10, by, cwid - 0.20, 0.6,
            [[('Dono: ', True), (p['dono'], False)],
             [('Prazo: ', True), (p['prazo'], False)]],
            size=9.5, color=GREY1, lsp=1.25, space_before=0.05)
        if i < 3:
            shp(sl, MSO_SHAPE.CHEVRON, x + cwid + 0.06, cy + 0.58, 0.22, 0.26,
                GREY4, None, adj=(0.50,))

    fy = 6.46
    rect(sl, M, fy, CW, 0.48, RED_BG, None)
    icon(sl, 'trophy', M + 0.16, fy + 0.09, 0.30)
    txt(sl, M + 0.62, fy, CW - 0.76, 0.48,
        [[(MERITO[0] + MERITO[1], True)]],
        size=10.5, color=RED, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    footer(sl, SRC_LQ, 1, logo=True)
    return sl


# ------------------------------------------------------------------ pagina 5
def page_fi_visao_v2(prs):
    sl = blank(prs)
    txt(sl, M, 0.24, CW, 0.2, [[('F&I MULTIBANCO', False)]], size=T_EYE,
        color=RED, spc=170)
    hline(sl, M, 0.48, CW, RGBColor(0xE9, 0xB8, 0xB8), 0.012)
    title(sl, *TIT_FI, y=0.72)

    txt(sl, M, 1.52, 4.0, 0.26, [[('Ambição', True)]], size=11.5, color=RED,
        bold=True)
    ay, ah = 1.82, 0.72
    rect(sl, M, ay, CW, ah, BAND2, None)
    icon(sl, 'target', M + 0.24, ay + 0.16, 0.40)
    txt(sl, M + 0.88, ay, CW - 1.10, ah, [[(AMBICAO, True)]], size=12,
        color=GREY1, bold=True, anchor=MSO_ANCHOR.MIDDLE, lsp=1.22)

    vy, vh = 2.76, 1.16
    vgap = 0.20
    vw = (CW - 2 * vgap) / 3.0
    for i, v in enumerate(VALOR):
        x = M + i * (vw + vgap)
        rect(sl, x, vy, vw, vh, CARD, None)
        icon(sl, v['ic'], x + 0.22, vy + 0.16, 0.32)
        txt(sl, x + 0.66, vy + 0.16, vw - 0.86, 0.30, [[(v['tit'], True)]],
            size=11.5, color=BLACK, bold=True)
        txt(sl, x + 0.22, vy + 0.60, vw - 0.44, 0.50, [[(v['txt'], False)]],
            size=10, color=GREY1, lsp=1.20)

    txt(sl, M, 4.08, 6.0, 0.26,
        [[('Como: ', True), ('cinco frentes, escaladas em ondas', True)]],
        size=11.5, color=RED, bold=True)

    fy, fgap = 4.44, 0.18
    fw = (CW - 4 * fgap) / 5.0
    for i, f in enumerate(FRENTES):
        x = M + i * (fw + fgap)
        rect(sl, x, fy, fw, 0.46, BAND2, None)
        numbubble(sl, x + 0.10, fy + 0.07, 0.32, f['n'], size=11.5)
        txt(sl, x + 0.50, fy, fw - 0.60, 0.46,
            [[(f['tit'].replace('\n', ' '), True)]], size=11, color=BLACK,
            bold=True, anchor=MSO_ANCHOR.MIDDLE, lsp=1.05)
        rect(sl, x, fy + 0.46, fw, 0.74, RGBColor(0xF7, 0xF7, 0xF7), None)
        icon(sl, f['ic'], x + fw / 2 - 0.19, fy + 0.62, 0.38)
        txt(sl, x + 0.02, fy + 1.36, fw - 0.06, 1.1, [[(f['txt'], False)]],
            size=9.5,
            color=GREY1, lsp=1.20)

    footer(sl, SRC_F1, 2, logo=True)
    return sl


# ------------------------------------------------------------------ main
def main(dst):
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    page_lead_quente(prs)
    page_fi_visao(prs)
    page_fi_plano(prs)
    page_lead_quente_v2(prs)
    page_fi_visao_v2(prs)
    cp = prs.core_properties
    cp.title = 'Lead quente e F&I multibanco'
    cp.author = 'ConsultingInnovation@bain.com'
    cp.last_modified_by = 'ConsultingInnovation@bain.com'
    prs.save(dst)
    print('ok %s (%d paginas)' % (dst, len(prs.slides._sldIdLst)))


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1
         else '../output/lead_quente_fi_multibanco.pptx')
