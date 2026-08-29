# -*- coding: utf-8 -*-
"""Tokens visuais, metricas de texto e helpers de tabela do painel Localiza.

Compartilhado por build.py (gera o deck do zero) e update_deck.py (edita o
deck ja formatado no template Bain).
"""
from copy import deepcopy
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import ImageFont

# ---------- tokens Localiza (do :root do painel) ----------
GREEN        = RGBColor(0x78, 0xDE, 0x1F)
GREEN_DEEP   = RGBColor(0x3F, 0x7C, 0x0E)
BLACK        = RGBColor(0x1A, 0x1A, 0x1A)
TEXT_DARK    = RGBColor(0x33, 0x33, 0x33)
TEXT_REG     = RGBColor(0x6B, 0x6B, 0x6B)
TEXT_HELPER  = RGBColor(0x75, 0x75, 0x75)
NAV_IDLE     = RGBColor(0x4A, 0x4A, 0x4A)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_X_LIGHT = RGBColor(0xF7, 0xF7, 0xF7)
GRAY_LIGHT   = RGBColor(0xEC, 0xEC, 0xEC)
GRAY_MEDIUM  = RGBColor(0xD5, 0xD5, 0xD5)
STROKE       = RGBColor(0xE2, 0xE2, 0xE2)
PLACEHOLDER  = RGBColor(0x82, 0x82, 0x82)

STATUS = {
    'verde':    ('Concluído',    RGBColor(0xE6, 0xF7, 0xEE), RGBColor(0x14, 0x6A, 0x3D)),
    'amarelo':  ('Em progresso', RGBColor(0xFF, 0xF8, 0xE1), RGBColor(0x6B, 0x57, 0x00)),
    'vermelho': ('Não iniciado', RGBColor(0xFC, 0xEA, 0xEA), RGBColor(0xA3, 0x00, 0x00)),
}

FONT_NAME = 'Arial'
LANG = 'pt-BR'
CELL_MX = 0.07

_TTF = {False: '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
        True:  '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf'}
_fc = {}


# ---------- metricas de texto ----------
def _f(bold):
    if bold not in _fc:
        _fc[bold] = ImageFont.truetype(_TTF[bold], 200)
    return _fc[bold]


def text_w(s, pt, bold=False):
    """largura do texto em polegadas"""
    return _f(bold).getlength(s) / 200.0 * pt / 72.0


def wrap(s, pt, width_in, bold=False):
    words, lines, cur = s.split(), [], ''
    for w in words:
        cand = (cur + ' ' + w).strip()
        if cur and text_w(cand, pt, bold) > width_in:
            lines.append(cur); cur = w
        else:
            cur = cand
    if cur:
        lines.append(cur)
    return lines or ['']


def nlines(s, pt, width_in, bold=False):
    return len(wrap(s, pt, width_in, bold))


def lh(pt):
    return pt * 1.22 / 72.0


def split_bold(text, phrases):
    """quebra o texto em (trecho, negrito?) a partir dos trechos a destacar"""
    spans, low = [], text.lower()
    for ph in phrases or []:
        i = low.find(ph.lower().strip())
        if i >= 0:
            spans.append((i, i + len(ph.strip())))
    if not spans:
        return [(text, False)]
    spans.sort()
    merged = [list(spans[0])]
    for a, b in spans[1:]:
        if a <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a, b])
    out, pos = [], 0
    for a, b in merged:
        if a > pos:
            out.append((text[pos:a], False))
        out.append((text[a:b], True))
        pos = b
    if pos < len(text):
        out.append((text[pos:], False))
    return [(t, b) for t, b in out if t]


# ---------- runs e shapes ----------
def no_shadow(shape):
    el = shape._element
    st = el.find(qn('p:style'))
    if st is not None:
        el.remove(st)
    spPr = el.spPr
    for e in spPr.findall(qn('a:effectLst')):
        spPr.remove(e)
    spPr.append(spPr.makeelement(qn('a:effectLst'), {}))


def style_run(r, text, size, color, bold=False, italic=False):
    r.text = text
    f = r.font
    f.name = FONT_NAME
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    # idioma pt-BR: evita que o PowerPoint trate o texto como outro idioma
    # (corretor automatico mexendo em acentuacao ao editar)
    r._r.get_or_add_rPr().set('lang', LANG)


# ---------- tabelas ----------
_TC_ORDER = ['lnL', 'lnR', 'lnT', 'lnB', 'lnTlToBr', 'lnBlToTr', 'cell3D',
             'noFill', 'solidFill', 'gradFill', 'blipFill', 'pattFill', 'grpFill',
             'headers', 'extLst']


def _tc_insert(tcPr, el):
    idx = _TC_ORDER.index(el.tag.split('}')[1])
    for child in tcPr:
        cname = child.tag.split('}')[1]
        if cname in _TC_ORDER and _TC_ORDER.index(cname) > idx:
            child.addprevious(el)
            return
    tcPr.append(el)


def cell_border(cell, edge, color=None, pt=0.75):
    """edge: L, R, T, B — color None = sem borda"""
    tcPr = cell._tc.get_or_add_tcPr()
    tag = qn('a:ln' + edge)
    for e in tcPr.findall(tag):
        tcPr.remove(e)
    ln = tcPr.makeelement(tag, {'w': str(int(pt * 12700)), 'cap': 'flat',
                                'cmpd': 'sng', 'algn': 'ctr'})
    if color is None:
        ln.append(ln.makeelement(qn('a:noFill'), {}))
    else:
        fill = ln.makeelement(qn('a:solidFill'), {})
        fill.append(ln.makeelement(qn('a:srgbClr'), {'val': str(color)}))
        ln.append(fill)
    _tc_insert(tcPr, ln)


def cell_box(cell, fill=None, borders=(None, None, None, None), lw=0.75):
    """borders na ordem L, R, T, B"""
    for edge, col in zip(('L', 'R', 'T', 'B'), borders):
        cell_border(cell, edge, col, lw)
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid(); cell.fill.fore_color.rgb = fill


def cell_text(cell, parts, size, color, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP, bold=False, mt=0.05, mb=0.05, mx=CELL_MX):
    """parts: str ou lista de (trecho, negrito?)"""
    cell.margin_left = Inches(mx); cell.margin_right = Inches(mx)
    cell.margin_top = Inches(mt); cell.margin_bottom = Inches(mb)
    cell.vertical_anchor = anchor
    tf = cell.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    if isinstance(parts, str):
        parts = [(parts, bold)]
    for seg, b in parts:
        style_run(p.add_run(), seg, size, color, bold or b)
    return tf


def set_table_plain(tbl):
    """desliga banding e o estilo padrao (azul) da tabela"""
    tbl.first_row = False
    tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr
    for e in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(e)
    sid = tblPr.makeelement(qn('a:tableStyleId'), {})
    sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'   # No Style, No Grid
    tblPr.append(sid)


def set_body_rows(tbl, n):
    """ajusta a quantidade de linhas de corpo (mantem a linha de cabecalho)"""
    trs = tbl._tbl.findall(qn('a:tr'))
    head, body = trs[0], trs[1:]
    while len(body) > n:
        tr = body.pop()
        tr.getparent().remove(tr)
    while len(body) < n:
        tr = deepcopy(body[-1] if body else head)
        head.getparent().append(tr)
        body.append(tr)


def set_row_heights(tbl, heights):
    for i, h in enumerate(heights):
        tbl.rows[i].height = Emu(int(round(h * 914400)))
